"""Tests of the header search: title rows, second header rows, comment lines and footers."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from sdg.audit import find_leaks, sensitive_values
from sdg.compare import compare
from sdg.generate import generate
from sdg.glimpse import glimpse
from sdg.matrix import read_matrix
from sdg.privacy import Options
from sdg.profile import profile, to_yaml
from sdg.report import Report, fact_map, render
from sdg.tables import data_end, find_header, read_tables, table_delimiter
from sdg.validate import validate


def write_workbook(path: Path, rows: list[list[object]]) -> Path:
    """Write a one-sheet workbook with the rows as given.

    Args:
        path: The file.
        rows: Every row, the rows above the header included.

    Returns:
        The path.
    """
    workbook = openpyxl.Workbook()
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)
    return path


def round_trip(path: Path, tmp_path: Path, options: Options) -> tuple[Report, Report, dict]:
    """Glimpse, profile and generate a file, then glimpse the synthetic file.

    Args:
        path: The real file.
        tmp_path: The pytest temporary folder.
        options: The glimpse and profile options.

    Returns:
        The real report, the synthetic report and the profiled dataset.
    """
    result = profile([path], options)
    config = tmp_path / "config.yaml"
    config.write_text(to_yaml(result), encoding="utf-8")
    out = tmp_path / "out"
    generate(config, out)
    dataset = result.config["datasets"][0]
    assert validate(out).issues == []
    return glimpse([path], options), glimpse([out / dataset["file"]], options), dataset


def errors(real: Report, synthetic: Report) -> list[str]:
    """List the errors of a compare.

    Args:
        real: The real report.
        synthetic: The synthetic report.

    Returns:
        The error messages.
    """
    return [item.message for item in compare(real, synthetic) if item.level == "error"]


@pytest.mark.parametrize(("records", "expected"), [
    ([["Study ABC-123 export"], [], ["id", "visit", "value"], ["S1", "V1", "1.5"]], 2),
    ([[None, "X1", "X2", "X3"], ["ID", "LIMIT_BAL", "SEX", "AGE"], [1, 20000, 2, 24]], 1),
    ([[None, "name", "type"], ["1", "TCGA-OR-A5J1", "ACC"]], 0),
    ([["Weblink", "https://gdc.cancer.gov"], [], ["TSS Code", "Source Site", "Disease"], ["01", "Site", "Ovary"],
      ["A1", "Site", "Lung"], ["B2", "Site", "Skin"], ["C3", "Site", "Ovary"]], 0),
    ([["Weblink", "https://gdc.cancer.gov"], [], ["TSS Code", "Source Site", "Disease"], ["1", "Site", "Ovary"],
      ["2", "Site", "Lung"]], 2),
    ([["#id", "name"], ["1", "a"], ["2", "b"]], 0),
    ([["", "Group", "Tissue"], ["S1", "A", "liver"], ["S2", "B", "lung"]], 0),
    ([["Sample", "Group", "Value"], ["S1", "A", ""], ["S2", "A", "3.4"]], 0),
    ([["Sample", "Group", "Value"], ["S1", "A", "<LOD"], ["S2", "A", "3.4"]], 0),
    ([["1", "2", "3"], ["4", "5", "6"]], 0),
    ([["Run date", "2026-10-01"], ["Operator", "J. Doe"], ["s", "a", "b", "c", "d"], ["S1", "1", "2", "3", "4"]], 2),
])
def test_find_header(records: list[list[object]], expected: int) -> None:
    """The header search skips title rows and a row of codes, and keeps row 1 when it finds no better row.

    The cases are a title and a blank row, the UCI credit codes row, the TCGA-CDR
    sheet with an unnamed first column, the TCGA-CDR TSS sheet with a link row
    and a text column of codes, the same sheet with a numeric column, a header
    that starts with #, a table of text with an unnamed column, a data row with
    an empty value, a data row with a missing token, a table without a header
    and an instrument block above the table.
    """
    assert find_header(records) == expected


def test_data_end_keeps_a_column_of_marked_values() -> None:
    """A GEO end line is a footer; a column of #12 values and a column that ends with #N/A are data."""
    geo = [["ID_REF", "GSM1"], ["p1", "1.0"], ["p2", "2.0"], ["!series_matrix_table_end"]]
    assert data_end(geo, 1) == 3
    marked = [["well", "value"], ["#1", "2"], ["#2", "3"]]
    assert data_end(marked, 1) == 3
    missing = [["days"], ["120"], ["#N/A"], ["#N/A"]]
    assert data_end(missing, 1) == 4


def test_delimiter_skips_comment_lines_and_decimal_commas() -> None:
    """The featureCounts command line holds a semicolon; a European CSV holds decimal commas."""
    feature_counts = ['# Program:featureCounts v2.0.1; Command:"featureCounts" "-a" "genes.gtf"',
                      "Geneid\tChr\tLength\ts1", "g1\tchr1;chr1;chr1\t1200\t5", "g2\tchr2\t900\t0"]
    assert table_delimiter(feature_counts, ".txt") == "\t"
    european = ["sample;value;ratio", "S1;1,5;0,25", "S2;2;0,5", "S3;3,25;1"]
    assert table_delimiter(european, ".csv") == ";"
    assert table_delimiter(["Report title, version 2", "a;b;c", "1;2;3", "4;5;6"], ".csv") == ";"


def test_title_rows_stay_out_of_the_report_and_the_config(tmp_path: Path) -> None:
    """A title and a blank row above the header: the columns read right, and the title never prints."""
    rows = [["Study ABC-123 export for Jane Doe"], [], ["subject_id", "visit", "value"]]
    rows += [[f"SUBJ-{index:03d}", f"visit_{index % 3}", 1.5 + index] for index in range(30)]
    path = write_workbook(tmp_path / "results.xlsx", rows)
    real, synthetic, dataset = round_trip(path, tmp_path, Options())
    facts = fact_map(real.tables[0].facts)
    assert facts["header row"] == "3" and facts["rows"] == "30"
    assert [column.name for column in real.tables[0].columns] == ["subject_id", "visit", "value"]
    assert dataset["sheets"][0]["preamble"] == [{"cells": 1}, {"cells": 0}]
    assert "Jane" not in render(real) and "Jane" not in to_yaml(profile([path], Options()))
    assert errors(real, synthetic) == []
    assert read_tables(tmp_path / "out" / dataset["file"])[0].preamble[0][0] == "note"


def test_second_header_row_makes_the_table_private(tmp_path: Path) -> None:
    """The UCI credit layout: codes X1 to X3 above the names. SEX and AGE make the table private."""
    rows = [[None, "X1", "X2", "X3"], ["ID", "LIMIT_BAL", "SEX", "AGE"]]
    rows += [[index, 10000 * (index % 7 + 1), index % 2 + 1, 21 + index % 50] for index in range(1, 41)]
    path = write_workbook(tmp_path / "credit.xlsx", rows)
    real, synthetic, dataset = round_trip(path, tmp_path, Options())
    facts = fact_map(real.tables[0].facts)
    assert facts["header row"] == "2" and facts["private"] == "safe facts"
    assert facts["private columns"] == "'SEX', 'AGE'"
    assert [column.name for column in real.tables[0].columns] == ["ID", "LIMIT_BAL", "SEX", "AGE"]
    assert set(fact_map(real.tables[0].columns[3].facts)) <= {"filled", "empty", "unique", "cell types",
                                                              "non-numeric"}
    assert dataset["sheets"][0]["preamble"] == [{"cells": 3, "start": 1}]
    assert errors(real, synthetic) == []


def test_names_in_the_first_data_row_make_the_table_private(tmp_path: Path) -> None:
    """When a full row of codes wins the header search, the names below it still mark the table."""
    rows = [["X0", "X1", "X2"], ["ID", "LIMIT_BAL", "AGE"]]
    rows += [[index, 10000 * (index % 7 + 1), 21 + index % 50] for index in range(1, 41)]
    path = write_workbook(tmp_path / "credit.xlsx", rows)
    report = glimpse([path], Options())
    assert fact_map(report.tables[0].facts)["private columns"] == "'X2'"
    age = fact_map(report.tables[0].columns[2].facts)
    assert "range" not in age and "5% to 95%" not in age and "codes" not in age


def test_comment_lines_above_the_header(tmp_path: Path) -> None:
    """A cBioPortal clinical file: four comment lines, then the header. The comments never print."""
    lines = ["#Patient Identifier\tSubtype\tStage", "#Identifier to uniquely specify a patient.\tSubtype\tStage",
             "#STRING\tSTRING\tSTRING", "#1\t1\t1", "PATIENT_ID\tSUBTYPE\tSTAGE"]
    lines += [f"TCGA-{index:04d}\tBRCA_{'LumA' if index % 2 else 'Basal'}\tSTAGE {index % 3 + 1}"
              for index in range(30)]
    path = tmp_path / "patients.txt"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    real, synthetic, dataset = round_trip(path, tmp_path, Options())
    facts = fact_map(real.tables[0].facts)
    assert facts["rows"] == "30" and facts["header row"] == "5" and facts["comment marks"] == "'#'"
    assert dataset["sheets"][0]["preamble"][0] == {"cells": 3, "mark": "#"}
    assert errors(real, synthetic) == []
    values = sensitive_values([path], Options())
    assert find_leaks(render(real), values) == []
    leak = find_leaks("description: Identifier to uniquely specify a patient.\n", values)
    assert leak[0].kind == "note"


def test_geo_series_matrix(tmp_path: Path) -> None:
    """A GEO series matrix: metadata lines, the table and the end line. The samples read right."""
    names = [f"GSM25{index:04d}" for index in range(12)]
    lines = ['!Series_title\t"Lung tumor study of Maria Landi"', "", "!Sample_title\t" + "\t".join(names),
             "!series_matrix_table_begin", '"ID_REF"\t' + "\t".join(f'"{name}"' for name in names)]
    lines += [f'"{probe}_at"\t' + "\t".join(f"{6 + probe * 0.1 + index * 0.01:.3f}" for index in range(12))
              for probe in range(20)]
    lines.append("!series_matrix_table_end")
    path = tmp_path / "GSE1_series_matrix.txt"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    options = Options(format="microarray")
    real, synthetic, dataset = round_trip(path, tmp_path, options)
    facts = fact_map(real.tables[0].facts)
    assert facts["samples"] == "12" and facts["features"] == "20"
    assert facts["header row"] == "5" and facts["footer rows"] == "1" and facts["comment marks"] == "'!'"
    assert dataset["footer"] == [{"cells": 1, "mark": "!"}]
    assert errors(real, synthetic) == []
    assert "Landi" not in render(real)
    assert find_leaks("title: Lung tumor study of Maria Landi\n", sensitive_values([path], options))


def test_featurecounts_command_line(tmp_path: Path) -> None:
    """The featureCounts command line above the header holds a semicolon; the matrix still reads by tab."""
    samples = [f"S{index:02d}.bam" for index in range(12)]
    lines = ['# Program:featureCounts v2.0.1; Command:"featureCounts" "-a" "genes.gtf"',
             "Geneid\t" + "\t".join(samples)]
    lines += [f"gene_{gene}\t" + "\t".join(str((gene * 7 + index * 3) % 50) for index in range(12))
              for gene in range(1, 41)]
    path = tmp_path / "counts.txt"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    matrix = read_matrix(path)
    assert matrix.delimiter == "\t" and len(matrix.samples) == 12 and len(matrix.features) == 40
    options = Options(format="counts")
    real, synthetic, dataset = round_trip(path, tmp_path, options)
    assert dataset["preamble"] == [{"cells": 1, "mark": "#"}]
    assert errors(real, synthetic) == []


def test_sheet_of_notes_reads_the_same_in_the_synthetic_file(tmp_path: Path) -> None:
    """The TCGA-CDR notes sheets: one note per row, so no row is wide. The synthetic file keeps row 1 as header.

    The synthetic cells of the two note columns are empty at random positions,
    so some synthetic rows fill both columns. Only a column of numbers, dates or
    booleans moves the header, and a sheet of notes has none.
    """
    rows: list[list[object]] = [["Recommended use of the endpoints:", None, None]]
    rows += [[None, f"Endpoint note number {index} on the use of PFI", None] if index % 2 else
             [None, None, f"Second note {index} on the censored cases"] for index in range(30)]
    path = write_workbook(tmp_path / "notes.xlsx", rows)
    real, synthetic, dataset = round_trip(path, tmp_path, Options())
    assert "preamble" not in dataset["sheets"][0]
    assert errors(real, synthetic) == []
