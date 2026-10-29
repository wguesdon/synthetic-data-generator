"""Tests of ``sdg relabel`` and of ``relabel_frame``, which apply a label key to real data."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from sdg import describe
from sdg.cli import main
from sdg.glimpse import glimpse
from sdg.privacy import LabelKey, Options
from sdg.relabel import relabel_frame
from sdg.report import Report, fact_map, parse, render


def generic_glimpse(targets: list[Path], key: Path) -> str:
    """Glimpse real files with generic labels and write their key.

    Args:
        targets: The real files or folders.
        key: The key file to write.

    Returns:
        The report text.
    """
    options = Options(generic_labels=True, key=LabelKey())
    text = render(glimpse(targets, options))
    options.key.write(key)
    return text


def test_copies_read_as_the_generic_glimpse(delivery: Path, tmp_path: Path) -> None:
    """A plain glimpse of each copy of a table that is not private equals the generic glimpse of the real file."""
    key = tmp_path / "key.csv"
    real = parse(generic_glimpse([delivery], key))
    assert main(["relabel", str(key), str(delivery), "-o", str(tmp_path / "copies")]) == 0
    copies = sorted((tmp_path / "copies").iterdir())
    assert [path.name for path in copies] == ["file_1.xlsx", "file_2.xlsx", "file_3.csv"]
    copied = parse(render(glimpse(copies, Options())))
    shared = [index for index, table in enumerate(real.tables) if "private" not in fact_map(table.facts)]
    assert shared == [0, 2]
    for index in shared:
        assert render(Report(tables=[copied.tables[index]])) == render(Report(tables=[real.tables[index]]))


def test_private_table_gets_generic_levels(delivery: Path, tmp_path: Path) -> None:
    """A private table gets generic column names and levels, as its report prints them."""
    key = tmp_path / "key.csv"
    generic_glimpse([delivery / "clinical.xlsx"], key)
    assert main(["relabel", str(key), str(delivery / "clinical.xlsx"), "-o", str(tmp_path / "copies")]) == 0
    workbook = openpyxl.load_workbook(tmp_path / "copies" / "file_1.xlsx")
    rows = list(workbook["sheet_1"].iter_rows(values_only=True))
    assert rows[0] == ("identifier_1", "numeric_1", "categorical_1", "categorical_2")
    assert {row[3] for row in rows[1:]} == {"level_1", "level_2"}


def test_workbook_sheets_swap_names_and_keep_the_title_rows(tmp_path: Path) -> None:
    """A real sheet called ``sheet_2`` becomes ``sheet_1``, and the rows above each header stay."""
    path = tmp_path / "trial.xlsx"
    workbook = openpyxl.Workbook()
    for index, title in enumerate(["sheet_2", "Data"]):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = title
        sheet.append([f"Trial export of {title}"])
        sheet.append([])
        sheet.append(["Patient", "Arm", "IL6"])
        for row in range(24):
            sheet.append([f"P{row:03d}", "drug" if row % 3 else "placebo", 1.5 + row * 0.25 + index])
    workbook.save(path)
    key = tmp_path / "key.csv"
    generic_glimpse([path], key)
    assert main(["relabel", str(key), str(path), "-o", str(tmp_path / "copies")]) == 0
    copy = openpyxl.load_workbook(tmp_path / "copies" / "file_1.xlsx")
    assert copy.sheetnames == ["sheet_1", "sheet_2"]
    rows = list(copy["sheet_2"].iter_rows(values_only=True))
    assert rows[0][0] == "Trial export of Data"
    assert rows[2] == ("identifier_1", "categorical_1", "numeric_1")
    assert rows[3] == ("P000", "level_2", 2.5) and rows[4] == ("P001", "level_1", 2.75)


def test_text_file_keeps_its_comment_lines(tmp_path: Path) -> None:
    """The comment lines above a CSV header stay, and the delimiter and the trailing spaces of levels hold."""
    path = tmp_path / "export.csv"
    lines = ["# exported by the plate reader", "Sample;Treatment;IL6"]
    lines += [f"S{index:02d};{'mid ' if index % 2 else 'low'};{1.5 + index:.2f}" for index in range(12)]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    key = tmp_path / "key.csv"
    generic_glimpse([path], key)
    assert main(["relabel", str(key), str(path), "-o", str(tmp_path / "copies")]) == 0
    copied = (tmp_path / "copies" / "file_1.csv").read_text(encoding="utf-8").splitlines()
    assert copied[:4] == ["# exported by the plate reader", "identifier_1;categorical_1;numeric_1",
                          "S00;level_1;1.50", "S01;level_2;2.50"]


def test_relabel_frame_and_back(tmp_path: Path) -> None:
    """A DataFrame takes the names of its glimpse, and ``reverse`` gives the real names back."""
    frame = pd.DataFrame({"Patient": [f"P{index:02d}" for index in range(12)],
                          "disease_status": ["patient", "healthy", "patient"] * 4,
                          "CD8": [10.0 + index for index in range(12)]})
    frame.loc[3, "disease_status"] = None
    key = tmp_path / "key.csv"
    report = parse(describe.glimpse(frame, name="panel.csv", generic_labels=True, label_key=key))
    renamed = relabel_frame(frame, key, "panel.csv")
    assert list(renamed.columns) == [column.name for column in report.tables[0].columns]
    assert renamed["categorical_1"].tolist()[:3] == ["level_1", "level_2", "level_1"]
    assert pd.isna(renamed["categorical_1"][3]) and renamed["categorical_1"].dtype == frame["disease_status"].dtype
    assert relabel_frame(frame, key, "file_1.csv").equals(renamed)
    pd.testing.assert_frame_equal(relabel_frame(renamed, key, "panel.csv", reverse=True), frame)


def test_relabel_frame_needs_a_sheet_of_a_workbook(tmp_path: Path) -> None:
    """With two sheets in the key, a DataFrame names its sheet by its real or its generic name."""
    path = tmp_path / "trial.xlsx"
    workbook = openpyxl.Workbook()
    for index, title in enumerate(["Notes", "Results"]):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = title
        sheet.append(["Patient", "IL6"])
        for row in range(12):
            sheet.append([f"P{row:02d}", 1.0 + row])
    workbook.save(path)
    key = tmp_path / "key.csv"
    generic_glimpse([path], key)
    frame = pd.read_excel(path, sheet_name="Results")
    with pytest.raises(ValueError, match="names 2 sheets"):
        relabel_frame(frame, key, path)
    assert list(relabel_frame(frame, key, path, sheet="Results").columns) == ["identifier_1", "numeric_1"]
    assert list(relabel_frame(frame, key, path, sheet="sheet_2").columns) == ["identifier_1", "numeric_1"]
    with pytest.raises(ValueError, match="no sheet 'Summary'"):
        relabel_frame(frame, key, path, sheet="Summary")


def test_cli_relabel_errors(delivery: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    """A file outside the key is skipped, the key is no target, a copy is never overwritten, and a key is checked."""
    key = delivery / "key.csv"
    generic_glimpse([delivery / "plate.csv"], key)
    output = tmp_path / "copies"
    assert main(["relabel", str(key), str(delivery), "-o", str(output)]) == 0
    err = capsys.readouterr().err
    assert "Skipped assay.xlsx: not in the label key" in err and "key.csv" not in err
    assert [path.name for path in output.iterdir()] == ["file_1.csv"]
    assert main(["relabel", str(key), str(delivery / "plate.csv"), "-o", str(output)]) == 1
    assert "exists; choose a new output folder" in capsys.readouterr().err
    assert main(["relabel", str(delivery / "plate.csv"), str(delivery / "plate.csv"), "-o", str(tmp_path / "x")]) == 1
    assert "not a label key" in capsys.readouterr().err


def test_unnamed_header_cell_in_a_merged_range(tmp_path: Path) -> None:
    """A header cell with no name inside a merged title gets its generic name, as the glimpse prints it."""
    path = tmp_path / "figure.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Figure 2"
    sheet.append(["Patient group"])
    sheet.append([None, "Arm", "IL6"])
    for row in range(24):
        sheet.append([f"P{row:03d}", "drug" if row % 3 else "placebo", 1.5 + row * 0.25])
    sheet.merge_cells("A1:A2")
    workbook.save(path)
    key = tmp_path / "key.csv"
    real = parse(generic_glimpse([path], key))
    assert main(["relabel", str(key), str(path), "-o", str(tmp_path / "copies")]) == 0
    copied = parse(render(glimpse([tmp_path / "copies" / "file_1.xlsx"], Options())))
    assert render(Report(tables=copied.tables)) == render(Report(tables=real.tables))
