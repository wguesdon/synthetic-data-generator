"""Tests of the generic column names by kind and of the label key (rule 3)."""

from __future__ import annotations

import csv
from pathlib import Path

import openpyxl
import pytest
import yaml

from conftest import SUBJECTS
from sdg.audit import find_leaks, sensitive_values
from sdg.cli import main
from sdg.compare import compare
from sdg.generate import generate
from sdg.glimpse import glimpse
from sdg.privacy import KEY_FIELDS, LabelKey, Options, kind_names
from sdg.profile import profile, to_yaml
from sdg.report import fact_map, render
from sdg.tables import write_text_table
from sdg.templates import init_config

#: The invented columns: two markers and a disease status.
MARKER_HEADER = ["Sample ID", "CD8", "CD4", "disease_status"]


def write_markers(path: Path) -> Path:
    """Write an invented table of marker frequencies and a disease status.

    Args:
        path: The CSV file.

    Returns:
        The path.
    """
    rows = [[SUBJECTS[index], f"{20 + index * 1.5:.1f}", f"{40 + index * 0.5:.1f}",
             "patient" if index % 4 else "healthy"] for index in range(len(SUBJECTS))]
    write_text_table(path, MARKER_HEADER, rows, ",")
    return path


def read_key(path: Path) -> list[dict[str, str]]:
    """Read a label key file.

    Args:
        path: The key file.

    Returns:
        One dict per row.
    """
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def generate_config(text: str, tmp_path: Path, name: str) -> Path:
    """Write a config text and generate it.

    Args:
        text: The YAML config.
        tmp_path: The pytest temporary folder.
        name: Name of the config and of the output folder.

    Returns:
        The output folder.
    """
    path = tmp_path / f"{name}.yaml"
    path.write_text(text, encoding="utf-8")
    generate(path, tmp_path / name)
    return tmp_path / name


def test_kind_names_count_per_kind_and_keep_a_repeat() -> None:
    """Each kind counts on its own, and a repeated name keeps the kind of its first column."""
    names = kind_names([("CD8", "numeric"), ("CD4", "numeric"), ("disease_status", "categorical"),
                        ("CD8", "text"), ("Visit", "date")])
    assert names == {"CD8": "numeric_1", "CD4": "numeric_2", "disease_status": "categorical_1", "Visit": "date_1"}


def test_glimpse_names_columns_by_kind(tmp_path: Path) -> None:
    """With generic labels, a marker is numeric_<n> and a status is categorical_<n>, with generic levels."""
    path = write_markers(tmp_path / "markers.csv")
    report = glimpse([path], Options(generic_labels=True))
    columns = report.tables[0].columns
    assert [column.name for column in columns] == ["identifier_1", "numeric_1", "numeric_2", "categorical_1"]
    assert fact_map(columns[3].facts)["values"] == "'level_1' 12, 'level_2' 4"
    text = render(report)
    assert "CD8" not in text and "disease_status" not in text and "patient" not in text
    assert find_leaks(text, sensitive_values([path], Options(generic_labels=True))) == []


def test_profile_names_match_the_glimpse(delivery: Path) -> None:
    """The profile and the glimpse give each column the same generic name."""
    options = Options(generic_labels=True)
    config = profile([delivery], options).config
    plate = next(item for item in config["datasets"] if item["file"] == "file_3.csv")
    names = [column["name"] for column in plate["sheets"][0]["columns"]]
    assert names == ["identifier_1", "categorical_1", "numeric_1", "text_1", "boolean_1", "date_1"]
    report = glimpse([delivery], options)
    assert [column.name for column in report.tables[2].columns] == names
    assert find_leaks(to_yaml(profile([delivery], options)), sensitive_values([delivery], options)) == []


def test_generic_round_trip_has_no_error(delivery: Path, tmp_path: Path) -> None:
    """A synthetic file made from a generic profile reads with the same generic names."""
    options = Options(generic_labels=True)
    out = generate_config(to_yaml(profile([delivery], options)), tmp_path, "out")
    differences = compare(glimpse([delivery], options), glimpse([out], options))
    assert [item for item in differences if item.level == "error"] == []


def test_profile_writes_the_label_key(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    """The key maps each generic file, column and level name back to the real name."""
    path = write_markers(tmp_path / "markers.csv")
    key = tmp_path / "key.csv"
    config = tmp_path / "config.yaml"
    assert main(["profile", str(path), "--generic-labels", "--label-key", str(key), "-o", str(config)]) == 0
    assert "keep it on the secure machine" in capsys.readouterr().err
    rows = read_key(key)
    assert list(rows[0]) == list(KEY_FIELDS)
    found = {(row["item"], row["column"], row["generic"]): row["real"] for row in rows}
    assert found[("file", "", "file_1.csv")] == str(path)
    assert found[("column", "", "numeric_1")] == "CD8" and found[("column", "", "numeric_2")] == "CD4"
    assert found[("column", "", "categorical_1")] == "disease_status"
    assert found[("level", "categorical_1", "level_1")] == "patient"
    assert found[("level", "categorical_1", "level_2")] == "healthy"
    assert "CD8" not in config.read_text(encoding="utf-8")


def test_glimpse_key_names_the_sheets(tmp_path: Path) -> None:
    """A workbook key holds the sheet name, and its columns carry the generic sheet name."""
    path = tmp_path / "panel.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Patient cohort"
    sheet.append(["CD8", "CD4"])
    for index in range(12):
        sheet.append([10 + index, 30 + index])
    workbook.save(path)
    key = tmp_path / "key.csv"
    assert main(["glimpse", str(path), "--generic-labels", "--label-key", str(key)]) == 0
    rows = {(row["item"], row["table"], row["generic"]): row["real"] for row in read_key(key)}
    assert rows[("sheet", "", "sheet_1")] == "Patient cohort"
    assert rows[("column", "sheet_1", "numeric_1")] == "CD8" and rows[("column", "sheet_1", "numeric_2")] == "CD4"


def test_audit_refuses_the_label_key(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    """A key holds the real names, so the audit fails on it even without generic labels."""
    path = write_markers(tmp_path / "markers.csv")
    key = tmp_path / "key.csv"
    assert main(["glimpse", str(path), "--generic-labels", "--label-key", str(key)]) == 0
    capsys.readouterr()
    assert main(["audit", str(key), str(path)]) == 1
    assert "label key" in capsys.readouterr().out


def test_olink_keeps_role_names_and_hides_other_columns(tmp_path: Path) -> None:
    """A role column keeps its standard name; another column gets the generic name that the glimpse prints."""
    first = generate_config(init_config(["olink"]), tmp_path, "first") / "dataset_1.csv"
    options = Options(generic_labels=True)
    config = profile([first], options).config
    text = to_yaml(profile([first], options))
    assert "Normalization" not in text and "Intensity" not in text
    names = [column["name"] for column in config["datasets"][0]["columns"]]
    assert names[:2] == ["SampleID", "Index"] and names[10] == "NPX"
    assert names[11] == glimpse([first], options).tables[0].columns[11].name == "categorical_6"
    assert find_leaks(text, sensitive_values([first], options)) == []
    keyed = Options(generic_labels=True, key=LabelKey())
    report = glimpse([first], keyed).tables[0]
    assert [column.name for column in report.columns] == names
    assert fact_map(report.facts)["value column"] == "'NPX'"
    roles = {column["name"] for column in config["datasets"][0]["columns"] if column["kind"] == "olink"}
    keyed_names = {row[5] for row in keyed.key.rows if row[0] == "column"}
    assert {"SampleID", "NPX"} <= roles and not roles & keyed_names and "Normalization" in keyed_names


def test_h5ad_glimpse_hides_obs_levels(tmp_path: Path) -> None:
    """With generic labels, the obs levels of an AnnData file print as generic levels."""
    data = yaml.safe_load(init_config(["h5ad"]))
    data["datasets"][0]["obs"].append({"name": "disease", "kind": "categorical",
                                       "levels": {"patient": 0.6, "healthy": 0.4}})
    out = generate_config(yaml.safe_dump(data, sort_keys=False), tmp_path, "out")
    report = glimpse([out / "dataset_1.h5ad"], Options(generic_labels=True))
    text = render(report)
    assert "patient" not in text and "disease" not in text and "leiden" not in text
    obs = report.tables[0].columns
    assert [column.name for column in obs][-2:] == ["categorical_1", "categorical_2"]
    assert fact_map(obs[-1].facts)["values"].startswith("'level_1'")
