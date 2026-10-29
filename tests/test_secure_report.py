"""Tests of ``sdg report``, the joined and audited report of the secure machine."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from sdg.audit import find_leaks, sensitive_values
from sdg.cli import main
from sdg.compare import compare
from sdg.describe import describe_file
from sdg.generate import generate
from sdg.glimpse import glimpse
from sdg.privacy import Options
from sdg.report import Report, parse, render
from sdg.secure_report import secure_report

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def study(tmp_path: Path) -> list[Path]:
    """Generate a study table with groups and a workbook with none.

    Args:
        tmp_path: The pytest temporary folder.

    Returns:
        The CSV file of the model example and the workbook of the minimal example.
    """
    generate(ROOT / "example" / "configs" / "model_study.yaml", tmp_path / "model")
    generate(ROOT / "example" / "configs" / "minimal_tabular.yaml", tmp_path / "minimal")
    return [tmp_path / "model" / "synthetic_assay.csv", tmp_path / "minimal" / "synthetic_clinical.xlsx"]


def test_report_joins_the_glimpse_and_the_descriptions(study: list[Path]) -> None:
    """The report holds the glimpse, then a design and a group part, and says that the audit passed."""
    result = secure_report(study, Options())
    assert result.findings == []
    assert result.text.startswith("# sdg report ") and "# audit: passed\n" in result.text
    assert render(glimpse(study, Options())) in result.text
    formats = [table.format for table in parse(result.text).tables]
    assert formats == ["csv", "excel", "design", "design", "groups"]
    assert "# file_2.xlsx | sheet 'sheet_1': no two groups to compare" in result.text


def test_compare_reads_the_glimpse_part_only(study: list[Path]) -> None:
    """A joined report compares with a glimpse report as the glimpse does."""
    result = secure_report(study, Options())
    assert compare(parse(result.text), glimpse(study, Options())) == []


def test_report_notes_a_file_that_is_not_a_table(tmp_path: Path) -> None:
    """A file of another format gets its glimpse and a note, and no description."""
    generate(ROOT / "example" / "configs" / "nanostring_rcc.yaml", tmp_path / "rcc")
    files = sorted((tmp_path / "rcc").rglob("*.RCC"))
    text = secure_report(files[:1], Options()).text
    assert "# file_1.rcc: not a plain table, no description" in text
    assert [table.format for table in parse(text).tables] == ["nanostring"]


def test_report_is_not_written_when_the_audit_fails(study: list[Path], tmp_path: Path,
                                                    monkeypatch: pytest.MonkeyPatch,
                                                    capsys: pytest.CaptureFixture[str]) -> None:
    """A real identifier in the joined text stops the report, and no file is written."""
    subject = next(value for value, kind in sensitive_values(study, Options()).items() if kind == "identifier")
    leaky = Report(comments=[f"a description that names {subject}"])
    monkeypatch.setattr("sdg.secure_report.describe_tables", lambda *arguments: leaky)
    output = tmp_path / "reports"
    assert main(["report", *map(str, study), "-o", f"{output}/"]) == 1
    assert "Audit failed" in capsys.readouterr().err
    assert list(output.glob("report_*.txt")) == []


def test_report_writes_into_a_new_folder(study: list[Path], tmp_path: Path) -> None:
    """``-o folder/`` makes the folder and writes ``report_<timestamp>.txt`` into it."""
    output = tmp_path / "new" / "reports"
    assert main(["report", *map(str, study), "--generic-labels", "-o", f"{output}/"]) == 0
    written = next(output.glob("report_*.txt")).read_text(encoding="utf-8")
    assert "'numeric_1'" in written and find_leaks(written, sensitive_values(study, Options(generic_labels=True))) == []


def write_two_sheets(path: Path, first: str) -> Path:
    """Write a workbook with two sheets of one invented table each.

    Args:
        path: The workbook.
        first: The title of the first sheet.

    Returns:
        The path.
    """
    workbook = openpyxl.Workbook()
    for index, title in enumerate([first, "Results"]):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = title
        sheet.append(["Patient", "Arm", "IL6"])
        for row in range(24):
            sheet.append([f"P{row:03d}", "drug" if row % 2 else "placebo", 1.5 + row * 0.25])
    workbook.save(path)
    return path


def test_describe_prints_the_glimpse_name_of_a_second_sheet(tmp_path: Path) -> None:
    """With generic labels, the second sheet of a workbook is ``sheet_2``, as in the glimpse."""
    path = write_two_sheets(tmp_path / "trial.xlsx", "Notes")
    report = describe_file("design", path, {}, Options(generic_labels=True), sheet="Results")
    assert report.tables[0].label == "file_1.xlsx | sheet 'sheet_2'"


def test_audit_skips_a_real_name_of_generic_form(tmp_path: Path) -> None:
    """Rule 3: a real sheet called ``sheet_1`` is no leak; a real sheet called ``Results`` stays one."""
    path = write_two_sheets(tmp_path / "trial.xlsx", "sheet_1")
    options = Options(generic_labels=True)
    values = sensitive_values([path], options)
    assert "sheet_1" not in values and values.get("Results") == "sheet name"
    assert find_leaks(render(glimpse([path], options)), values) == []
    assert find_leaks("== file_1.xlsx | sheet 'Results' <excel>", values) != []


def test_report_words_are_no_leak_of_a_column_name(tmp_path: Path) -> None:
    """Rule 3: real columns called ``value``, ``group`` and ``time`` do not match the words of a generic report."""
    path = tmp_path / "labs.csv"
    rows = ["subject,group,time,value"]
    rows += [f"S{index // 2:02d},{'drug' if index % 4 < 2 else 'placebo'},{index % 2},{1.5 + index:.2f}"
             for index in range(48)]
    path.write_text("\n".join(rows) + "\n", encoding="utf-8")
    options = Options(generic_labels=True)
    values = sensitive_values([path], options)
    assert {"value", "group", "time"} <= set(values)
    result = secure_report([path], options)
    assert result.findings == [] and "group column: 'categorical_1'" in result.text
    assert [finding.value for finding in find_leaks(result.text + "$ 'value' <numeric>\n", values)] == ["value"]
    config = "datasets:\n  - name: dataset_1\n    columns:\n      - {name: group, kind: categorical}\n"
    assert find_leaks(config, values) == []
    assert [finding.value for finding in find_leaks(config.replace("dataset_1", "drug"), values)] == ["drug"]
