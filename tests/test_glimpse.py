"""Tests of ``sdg glimpse`` and its privacy rules on CSV and Excel files."""

from __future__ import annotations

import datetime as dt
import re
import zipfile
from pathlib import Path

import openpyxl
import pytest

from sdg.cli import main
from sdg.glimpse import glimpse
from sdg.privacy import Options
from sdg.report import ColumnReport, Report, TableReport, fact_map, parse, parse_counted, render
from sdg.tables import write_text_table


def write_csv(path: Path, header: list[str], rows: list[list[object]], delimiter: str = ",") -> Path:
    """Write a small CSV file for a test.

    Args:
        path: The file.
        header: Column names.
        rows: Data rows.
        delimiter: The delimiter.

    Returns:
        The path.
    """
    write_text_table(path, header, [[str(cell) for cell in row] for row in rows], delimiter)
    return path


def column(report: Report, name: str, table: int = 0) -> ColumnReport:
    """Find one column report by name.

    Args:
        report: The glimpse report.
        name: Column name.
        table: Table position.

    Returns:
        The column report.
    """
    return next(item for item in report.tables[table].columns if item.name == name)


def test_numeric_column_with_na_tokens_prints_no_value(tmp_path: Path) -> None:
    """26 NA tokens in 120 rows keep the column numeric, and its values do not print."""
    values = [f"{10 + index * 0.0137:.4f}" for index in range(94)] + ["NA"] * 26
    path = write_csv(tmp_path / "plate.csv", ["reading"], [[value] for value in values])
    report = glimpse([path], Options())
    reading = column(report, "reading")
    facts = fact_map(reading.facts)
    assert reading.kind == "numeric"
    assert "values" not in facts
    assert parse_counted(facts["tokens"])[0] == {"NA": 26}
    text = render(report)
    assert not any(value in text for value in values[:94])


def test_identifier_prints_shapes_only(tmp_path: Path) -> None:
    """Subject identifiers print as shapes, never as values."""
    ids = [f"ABC-{index:04d}" for index in range(1, 21)]
    path = write_csv(tmp_path / "ids.csv", ["Subject ID"], [[value] for value in ids])
    report = glimpse([path], Options())
    subject = column(report, "Subject ID")
    assert subject.kind == "identifier"
    assert parse_counted(fact_map(subject.facts)["shapes"])[0] == {"AAA-9999": 20}
    assert not any(value in render(report) for value in ids)


def test_private_column_prints_counts_only(tmp_path: Path) -> None:
    """Rule 2: an age column makes the table private; the age prints counts only, the visit generic levels."""
    rows = [[f"{20 + index}", "visit_1"] for index in range(30)]
    path = write_csv(tmp_path / "demo.csv", ["Age", "visit"], rows)
    report = glimpse([path], Options())
    facts = fact_map(report.tables[0].facts)
    assert facts["private"] == "safe facts" and facts["private columns"] == "'Age'"
    age, visit = report.tables[0].columns
    assert set(fact_map(age.facts)) <= {"filled", "empty", "unique", "cell types", "non-numeric"}
    assert fact_map(visit.facts)["values"] == "'level_1' 30" and "visit_1" not in render(report)


def test_private_file_name_prints_generic_levels(tmp_path: Path) -> None:
    """A file whose name matches the private pattern prints its levels as generic counts only."""
    path = write_csv(tmp_path / "clinical_export.csv", ["visit"], [["v1"], ["v1"], ["v2"], ["v2"]])
    report = glimpse([path], Options())
    assert fact_map(report.tables[0].columns[0].facts)["values"] == "'level_1' 2, 'level_2' 2"
    assert "v1" not in render(report) and "private columns" not in fact_map(report.tables[0].facts)


def test_few_numbers_print_no_range(tmp_path: Path) -> None:
    """Fewer than 10 numbers record counts only (rule 5)."""
    path = write_csv(tmp_path / "small.csv", ["x"], [[str(index)] for index in range(5)])
    facts = fact_map(column(glimpse([path], Options()), "x").facts)
    assert "range" not in facts and facts["numbers"] == "5"


def test_numbers_are_rounded_to_three_digits(tmp_path: Path) -> None:
    """Range and median keep three significant digits."""
    path = write_csv(tmp_path / "round.csv", ["x"], [[f"{1.23456 + index}"] for index in range(12)])
    facts = fact_map(column(glimpse([path], Options()), "x").facts)
    assert facts["range"] == "1.23 to 12.2"
    assert facts["decimals"] == "5"


def test_categorical_keeps_stray_space(tmp_path: Path) -> None:
    """A trailing space in a level survives."""
    rows = [["visit_1"]] * 10 + [["visit_2 "]] * 10
    path = write_csv(tmp_path / "visits.csv", ["visit"], rows)
    visit = column(glimpse([path], Options()), "visit")
    assert visit.kind == "categorical"
    assert parse_counted(fact_map(visit.facts)["values"])[0] == {"visit_1": 10, "visit_2 ": 10}


def test_free_text_prints_lengths_only(tmp_path: Path) -> None:
    """Unique comments are free text: lengths and shapes, no values."""
    comments = [f"note number {index} about the sample" for index in range(12)]
    path = write_csv(tmp_path / "notes.csv", ["comment"], [[value] for value in comments])
    report = glimpse([path], Options())
    facts = fact_map(column(report, "comment").facts)
    assert column(report, "comment").kind == "text"
    assert "values" not in facts and "lengths" in facts
    assert not any(value in render(report) for value in comments)


def test_table_facts(tmp_path: Path) -> None:
    """The report gives the delimiter, the repeated names and the empty rows."""
    path = write_csv(tmp_path / "semi.csv", ["a", "b", "a"], [["1", "x", "2"], ["", "", ""], ["3", "y", "4"]],
                     delimiter=";")
    facts = fact_map(glimpse([path], Options()).tables[0].facts)
    assert facts["delimiter"] == "';'"
    assert facts["repeated names"] == "'a'"
    assert facts["empty rows"] == "1"
    assert facts["rows"] == "3"


def test_workbook_types(tmp_path: Path) -> None:
    """A workbook keeps its cell types: text numbers, dates and booleans."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    sheet.append(["value", "when", "flag"])
    for index in range(12):
        sheet.append([str(index * 1.5), dt.datetime(2026, 1, 1) + dt.timedelta(days=index), index % 2 == 0])
    path = tmp_path / "book.xlsx"
    workbook.save(path)
    report = glimpse([path], Options())
    table = report.tables[0]
    assert table.label == "file_1.xlsx | sheet 'Results'" and table.format == "excel"
    assert fact_map(column(report, "value").facts)["cell types"] == "str 12"
    when = fact_map(column(report, "when").facts)
    assert column(report, "when").kind == "date" and when["span days"] == "11"
    assert fact_map(column(report, "flag").facts)["true"] == "6"
    assert "2026" not in render(report)


def break_dimension(path: Path) -> None:
    """Rewrite the dimension tag of the first sheet to ``A1:A1``, as some writers do.

    Args:
        path: The ``.xlsx`` file, changed in place.
    """
    with zipfile.ZipFile(path) as source:
        members = {name: source.read(name) for name in source.namelist()}
    sheet = "xl/worksheets/sheet1.xml"
    members[sheet] = re.sub(rb'<dimension ref="[^"]*"', b'<dimension ref="A1:A1"', members[sheet])
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, data in members.items():
            target.writestr(name, data)


def test_workbook_with_a_wrong_dimension_tag(tmp_path: Path) -> None:
    """A sheet whose dimension tag says A1:A1 still reads every cell."""
    workbook = openpyxl.Workbook()
    workbook.active.append(["value", "group"])
    for index in range(12):
        workbook.active.append([index * 1.5, "a" if index % 2 else "b"])
    path = tmp_path / "book.xlsx"
    workbook.save(path)
    break_dimension(path)
    table = glimpse([path], Options()).tables[0]
    assert fact_map(table.facts)["rows"] == "12" and [item.name for item in table.columns] == ["value", "group"]


def test_generic_labels(tmp_path: Path) -> None:
    """Generic labels hide column names, sheet names and levels."""
    rows = [["drug_x"]] * 6 + [["drug_y"]] * 6
    path = write_csv(tmp_path / "arms.csv", ["Treatment arm"], rows)
    report = glimpse([path], Options(generic_labels=True))
    text = render(report)
    assert "Treatment" not in text and "drug_x" not in text
    assert parse_counted(fact_map(report.tables[0].columns[0].facts)["values"])[0] == {"level_1": 6, "level_2": 6}


def test_file_names_never_print(tmp_path: Path) -> None:
    """Rule 8: the report labels files generically."""
    first = write_csv(tmp_path / "zz_first_patient_export.csv", ["x"], [["1"]])
    second = write_csv(tmp_path / "zz_second.csv", ["x"], [["1"]])
    text = render(glimpse([first, second], Options()))
    assert "zz_first" not in text and "zz_second" not in text
    assert "== file_1.csv <csv>" in text and "== file_2.csv <csv>" in text


def test_render_and_parse_round_trip(tmp_path: Path) -> None:
    """Parsing a rendered report gives the same report."""
    rows = [["visit; one"], ["visit; one"], ["it's"], ["it's"]]
    path = write_csv(tmp_path / "quotes.csv", ["na'me; x"], rows)
    report = glimpse([path], Options())
    assert parse(render(report)) == report


def test_parse_ignores_lines_around_a_pasted_report() -> None:
    """Text around a pasted report does not break the parser."""
    report = Report(tables=[TableReport("file_1.csv", "csv", columns=[ColumnReport("x", "numeric")])])
    pasted = "Here is the report:\n" + render(report) + "\nThanks"
    assert parse(pasted).tables == report.tables


def test_cli_writes_a_timestamped_report(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    """The -o option writes glimpse_<stamp>.txt into a directory."""
    path = write_csv(tmp_path / "data.csv", ["x"], [["1"]])
    out_dir = tmp_path / "reports"
    out_dir.mkdir()
    assert main(["glimpse", str(path), "-o", str(out_dir)]) == 0
    written = list(out_dir.glob("glimpse_*.txt"))
    assert len(written) == 1 and re.fullmatch(r"glimpse_\d{4}_\d{2}_\d{2}_\d{6}\.txt", written[0].name)
    assert "Wrote" in capsys.readouterr().out


def test_cli_missing_target(capsys: pytest.CaptureFixture[str]) -> None:
    """A missing target is an error, not a crash."""
    assert main(["glimpse", "does_not_exist.csv"]) == 1
    assert "not found" in capsys.readouterr().err


def test_numeric_column_with_mostly_missing_tokens_stays_numeric(tmp_path: Path) -> None:
    """A survival column of the TCGA-CDR workbook is 68% "#N/A"; the numeric share leaves the token out."""
    values = [str(100 + index * 37) for index in range(30)] + ["#N/A"] * 70
    path = write_csv(tmp_path / "survival.csv", ["death_days_to"], [[value] for value in values])
    days = column(glimpse([path], Options()), "death_days_to")
    facts = fact_map(days.facts)
    assert days.kind == "numeric" and "values" not in facts
    assert parse_counted(facts["tokens"])[0] == {"#N/A": 70}


def test_one_number_among_missing_tokens_never_prints(tmp_path: Path) -> None:
    """A lone number in a column of "#N/A" printed as a category level; it is now a count."""
    path = write_csv(tmp_path / "lone.csv", ["dfi_time"], [["#N/A"]] * 40 + [["8123"]])
    report = glimpse([path], Options())
    assert column(report, "dfi_time").kind == "numeric"
    assert "8123" not in render(report)


def test_identifier_name_with_a_qualifier_keeps_its_levels(tmp_path: Path) -> None:
    """Sample_Type and SampleQC of an Olink export hold levels that analysis code filters on."""
    rows = [[f"S{index:03d}", "CONTROL" if index % 6 == 0 else "SAMPLE", "PASS"] for index in range(30)]
    path = write_csv(tmp_path / "qc.csv", ["SampleID", "Sample_Type", "SampleQC"], rows)
    report = glimpse([path], Options())
    assert column(report, "SampleID").kind == "identifier"
    sample_type = column(report, "Sample_Type")
    assert sample_type.kind == "categorical"
    assert parse_counted(fact_map(sample_type.facts)["values"])[0] == {"SAMPLE": 25, "CONTROL": 5}
    assert column(report, "SampleQC").kind == "categorical"


def test_question_mark_is_a_missing_token(tmp_path: Path) -> None:
    """UCI and MIMIC files write ? for a missing value, so a column of numbers and ? is numeric."""
    values = [f"{index}.0" for index in range(30)] + ["?"] * 60
    path = write_csv(tmp_path / "risk.csv", ["Smokes (years)"], [[value] for value in values])
    years = column(glimpse([path], Options()), "Smokes (years)")
    assert years.kind == "numeric"
    assert parse_counted(fact_map(years.facts)["tokens"])[0] == {"?": 60}


def test_platelets_is_not_an_identifier(tmp_path: Path) -> None:
    """The name platelets holds plate, but its values are counts; PlateID stays an identifier."""
    rows = [[f"{200000 + index * 1371}", f"P{index % 3}"] for index in range(30)]
    path = write_csv(tmp_path / "blood.csv", ["platelets", "PlateID"], rows)
    report = glimpse([path], Options())
    assert column(report, "platelets").kind == "numeric"
    assert column(report, "PlateID").kind == "identifier"
