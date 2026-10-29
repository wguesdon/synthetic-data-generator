"""Tests of ``sdg generate`` in the resample mode, for CSV and Excel."""

from __future__ import annotations

import datetime as dt
import json
import re
import uuid
from collections import Counter
from pathlib import Path

import numpy as np
import openpyxl
import pytest
import yaml

from sdg.cli import main
from sdg.compare import compare
from sdg.config import ConfigError, parse_config
from sdg.formats.tabular import NumericColumn, TextColumn, column_cells
from sdg.modes.resample import (invented_identifiers, pattern_capacity, pattern_counts, pick_shares,
                                share_counts)
from sdg.generate import generate
from sdg.glimpse import glimpse
from sdg.audit import find_leaks, sensitive_values
from sdg.output import file_sha256
from sdg.privacy import Options, shape
from sdg.profile import profile, to_yaml
from sdg.report import fact_map
from sdg.tables import read_tables

EXAMPLE = Path(__file__).resolve().parents[1] / "example" / "configs" / "minimal_tabular.yaml"

#: A pattern that matches no file name, so no output table is private.
NO_PRIVATE_FILES = re.compile(r"(?!x)x")


def structure(config: dict) -> list[tuple]:
    """Reduce a profiled config to its structure, without the measured values.

    Args:
        config: A profiled config.

    Returns:
        For each dataset and sheet: the format, the delimiter, the sheet name,
        the row counts and each column name, kind and level set.
    """
    found = []
    for dataset in config["datasets"]:
        for sheet in dataset["sheets"]:
            columns = [(column["name"], column["kind"], tuple(sorted(column.get("levels", {}))))
                       for column in sheet["columns"]]
            found.append((dataset["format"], dataset.get("delimiter"), sheet["name"], sheet["rows"],
                          sheet["empty_rows"], columns))
    return found


def test_example_output(tmp_path: Path) -> None:
    """The example workbook has the configured sheet, header, rows and empty rows."""
    files = generate(EXAMPLE, tmp_path / "out")
    names = sorted(path.name for path in files)
    assert names == ["SYNTHETIC_DATA.md", "config.yaml", "manifest.json", "synthetic_clinical.xlsx"]
    table = read_tables(tmp_path / "out" / "synthetic_clinical.xlsx")[0]
    assert table.sheet == "sheet_1"
    assert table.header == ["subject_id", "visit", "feature_01", "comment", "sample_date", "responder"]
    assert len(table.rows) == 120
    assert sum(all(cell is None for cell in row) for row in table.rows) == 6


def test_glimpse_of_output_matches_config_kinds(tmp_path: Path) -> None:
    """Each generated column reads back as its configured kind."""
    generate(EXAMPLE, tmp_path / "out")
    report = glimpse([tmp_path / "out" / "synthetic_clinical.xlsx"], Options(private_files=NO_PRIVATE_FILES))
    kinds = [column.kind for column in report.tables[0].columns]
    assert kinds == ["identifier", "categorical", "numeric", "text", "date", "boolean"]
    assert fact_map(report.tables[0].facts)["empty rows"] == "6"


def test_same_seed_same_data(tmp_path: Path, delivery: Path) -> None:
    """Criterion 5: the same config and seed give the same data after read."""
    config = tmp_path / "config.yaml"
    config.write_text(to_yaml(profile([delivery], Options(), seed=3)), encoding="utf-8")
    generate(config, tmp_path / "one")
    generate(config, tmp_path / "two")
    assert file_sha256(tmp_path / "one" / "file_3.csv") == file_sha256(tmp_path / "two" / "file_3.csv")
    for name in ("file_1.xlsx", "file_2.xlsx"):
        first = read_tables(tmp_path / "one" / name)
        second = read_tables(tmp_path / "two" / name)
        assert [table.rows for table in first] == [table.rows for table in second]


def test_profile_of_output_has_the_input_structure(tmp_path: Path, delivery: Path) -> None:
    """Profile, generate, profile again: the structure survives the round trip."""
    first = profile([delivery], Options()).config
    config = tmp_path / "config.yaml"
    config.write_text(yaml.safe_dump(first, sort_keys=False), encoding="utf-8")
    generate(config, tmp_path / "out")
    outputs = [tmp_path / "out" / dataset["file"] for dataset in first["datasets"]]
    second = profile(outputs, Options()).config
    assert structure(second) == structure(first)


def test_shares_are_close_to_the_config(tmp_path: Path) -> None:
    """Missing shares, level shares and the true share follow the config."""
    data = yaml.safe_load(EXAMPLE.read_text(encoding="utf-8"))
    sheet = data["datasets"][0]["sheets"][0]
    sheet["rows"], sheet["empty_rows"] = 4000, 0
    config = tmp_path / "big.yaml"
    config.write_text(yaml.safe_dump(data), encoding="utf-8")
    generate(config, tmp_path / "out")
    table = read_tables(tmp_path / "out" / "synthetic_clinical.xlsx")[0]
    comment = table.column(3)
    assert abs(sum(cell is None for cell in comment) / 4000 - 0.8) < 0.03
    feature = table.column(2)
    assert abs(sum(cell == "NA" for cell in feature) / 4000 - 0.22) < 0.03
    assert abs(sum(cell is True for cell in table.column(5)) / 4000 - 0.35) < 0.03
    numbers = [cell for cell in feature if isinstance(cell, float)]
    assert min(numbers) >= 0.12 and max(numbers) <= 9.8


def test_identifiers_count_up_under_one_prefix() -> None:
    """Invented identifiers share their letters and count up, as real ones do."""
    values = invented_identifiers("AA-9999", 3, np.random.default_rng(0))
    assert len(set(value[:2] for value in values)) == 1
    assert [value[2:] for value in values] == ["-0001", "-0002", "-0003"]
    assert len(set(invented_identifiers("A9", 200, np.random.default_rng(1)))) == 200


def test_integer_identifiers_have_no_leading_zero() -> None:
    """A shape of digits only counts up from a drawn start within its width, with no leading zero."""
    values = invented_identifiers("99", 3, np.random.default_rng(0))
    assert [int(value) - int(values[0]) for value in values] == [0, 1, 2]
    assert all(len(value) == 2 and not value.startswith("0") for value in values)
    assert invented_identifiers("9", 9, np.random.default_rng(0)) == [str(digit) for digit in range(1, 10)]
    starts = {invented_identifiers("99999999", 100, np.random.default_rng(seed))[0] for seed in range(5)}
    assert len(starts) == 5 and "10000000" not in starts


def test_pattern_counts_respect_the_room_of_each_shape() -> None:
    """A shape never gets more identifiers than it can hold while another shape has room."""
    counts = pattern_counts({"99": 0.7, "999": 0.2, "9": 0.1}, 152)
    assert sum(counts) == 152 and counts[2] <= 9 and counts[0] <= 90
    assert pattern_capacity("A9") == 260 and pattern_capacity("999") == 900


def test_identifier_with_mixed_shapes_stays_integer(tmp_path: Path) -> None:
    """Sample numbers 1 to 152 keep their three shapes and stay distinct integers."""
    numbers = [str(index) for index in range(1, 153)] * 2
    path = tmp_path / "samples.csv"
    path.write_text("Sample Number\n" + "\n".join(numbers) + "\n", encoding="utf-8")
    result = profile([path], Options())
    column = result.config["datasets"][0]["sheets"][0]["columns"][0]
    assert set(column["pattern"]) == {"9", "99", "999"} and column["distinct"] == 152
    config = tmp_path / "config.yaml"
    config.write_text(to_yaml(result), encoding="utf-8")
    generate(config, tmp_path / "out")
    cells = [str(cell) for cell in read_tables(tmp_path / "out" / "file_1.csv")[0].column(0)]
    assert all(cell.isdigit() and not cell.startswith("0") for cell in cells)
    assert len({int(cell) for cell in cells}) == 152
    assert {len(cell) for cell in cells} == {1, 2, 3}


def test_refuses_a_folder_that_sdg_did_not_write(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    """A mistyped output path never overwrites other files."""
    folder = tmp_path / "real_data"
    folder.mkdir()
    (folder / "keep.csv").write_text("x\n1\n", encoding="utf-8")
    assert main(["generate", str(EXAMPLE), "-o", str(folder)]) == 1
    assert "did not write" in capsys.readouterr().err
    assert (folder / "keep.csv").read_text(encoding="utf-8") == "x\n1\n"


def test_reuses_an_earlier_output(tmp_path: Path) -> None:
    """A second run into the same sdg folder works."""
    generate(EXAMPLE, tmp_path / "out")
    assert generate(EXAMPLE, tmp_path / "out")


def test_manifest_matches_files(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    """The manifest lists every data file with its hash."""
    assert main(["generate", str(EXAMPLE), "-o", str(tmp_path / "out"), "--json"]) == 0
    listed = json.loads(capsys.readouterr().out)["files"]
    manifest = json.loads((tmp_path / "out" / "manifest.json").read_text(encoding="utf-8"))
    for entry in manifest["files"]:
        assert entry["sha256"] == file_sha256(tmp_path / "out" / entry["path"])
    assert {entry["path"] for entry in manifest["files"]} | {"manifest.json"} == set(listed)


def test_csv_text_quirks(tmp_path: Path) -> None:
    """Delimiter, na_token, boolean texts and the date format reach the CSV text."""
    config = parse_config({
        "version": 1, "seed": 1,
        "datasets": [{"name": "t", "format": "csv", "file": "t.csv", "delimiter": ";", "sheets": [{
            "rows": 50, "columns": [
                {"name": "v", "kind": "numeric", "quantiles": [1, 2], "decimals": 2, "missing": 0.5,
                 "na_token": "N/A"},
                {"name": "b", "kind": "boolean", "true_share": 0.5, "true_text": "yes", "false_text": "no"},
                {"name": "d", "kind": "date", "interval_days": [0, 10], "date_format": "%d/%m/%Y"}]}]}]})
    path = tmp_path / "config.yaml"
    path.write_text(yaml.safe_dump(config.model_dump(mode="json")), encoding="utf-8")
    generate(path, tmp_path / "out")
    lines = (tmp_path / "out" / "t.csv").read_text(encoding="utf-8").splitlines()
    assert lines[0] == "v;b;d"
    cells = [line.split(";") for line in lines[1:]]
    assert {cell[1] for cell in cells} == {"yes", "no"}
    assert "N/A" in {cell[0] for cell in cells}
    assert all(len(cell[0]) == 4 for cell in cells if cell[0] != "N/A")
    assert all(cell[2].endswith("/2000") for cell in cells)


def test_share_counts_are_exact() -> None:
    """Each share gets its rounded count of cells: one token in 17 cells is one cell for every seed."""
    assert share_counts([0.5, 0.25], 10).tolist() == [5, 3, 2]
    for seed in range(20):
        assert (pick_shares([1 / 17], 17, np.random.default_rng(seed)) == 0).sum() == 1


def test_tokens_keep_their_cells_beside_empty_cells() -> None:
    """An empty cell never takes the place of a token, so both keep their counts."""
    column = NumericColumn(name="p", kind="numeric", quantiles=[0.0, 1.0], tokens={"p_value": 1 / 21},
                           missing=4 / 21)
    for seed in range(20):
        cells = column_cells(column, 21, np.random.default_rng(seed))
        assert cells.count("p_value") == 1 and cells.count(None) == 4


def test_text_cells_are_never_empty() -> None:
    """A filled text cell has one character or more; only the missing share empties cells."""
    column = TextColumn(name="note", kind="text", lengths=(0, 3), missing=0.1)
    for seed in range(20):
        cells = column_cells(column, 50, np.random.default_rng(seed))
        assert cells.count(None) == 5 and all(cell for cell in cells if cell is not None)
    empty = TextColumn(name="note", kind="text", lengths=(0, 0))
    assert column_cells(empty, 3, np.random.default_rng(0)) == ["", "", ""]


def test_text_column_keeps_its_common_shapes(tmp_path: Path) -> None:
    """Rule 6: a text column records the shapes that the report prints, never a value.

    Cabin codes fill most cells with three shapes, so the profile keeps them and
    the synthetic cells follow them. Free notes have no common shape, so they
    keep the length range only.
    """
    cabins = [f"{'ABCDEFG'[index % 7]}{10 + index}" for index in range(40)] + ["E101", "F G73"] * 5
    notes = [f"note {'x' * index} seen" for index in range(50)]
    path = tmp_path / "cabins.csv"
    path.write_text("Cabin,remark\n" + "\n".join(f"{cabin},{note}" for cabin, note in zip(cabins, notes))
                    + "\n", encoding="utf-8")
    result = profile([path], Options())
    cabin, remark = result.config["datasets"][0]["sheets"][0]["columns"]
    assert cabin["kind"] == "text" and cabin["shapes"] == {"A99": 0.8, "A999": 0.1, "A A99": 0.1}
    assert "shapes" not in remark
    text = to_yaml(result)
    assert find_leaks(text, sensitive_values([path], Options())) == []
    config = tmp_path / "config.yaml"
    config.write_text(text, encoding="utf-8")
    generate(config, tmp_path / "out")
    cells = read_tables(tmp_path / "out" / "file_1.csv")[0].column(0)
    assert Counter(shape(cell) for cell in cells) == Counter({"A99": 40, "A999": 5, "A A99": 5})


def test_private_text_column_keeps_its_shapes(tmp_path: Path) -> None:
    """Rule 2: a text column of a private table keeps its common shapes, a safe fact, and the default lengths."""
    rows = [f"{40 + index},C{10 + index}" for index in range(30)]
    path = tmp_path / "demo.csv"
    path.write_text("Age,Cabin\n" + "\n".join(rows) + "\n", encoding="utf-8")
    cabin = profile([path], Options()).config["datasets"][0]["sheets"][0]["columns"][1]
    assert cabin == {"name": "Cabin", "kind": "text", "lengths": [1, 20], "shapes": {"A99": 1.0}}
    assert fact_map(glimpse([path], Options()).tables[0].columns[1].facts)["shapes"] == "'A99' 30"


def test_private_text_column_gets_lengths_from_one(tmp_path: Path) -> None:
    """The default lengths of a private text column start at 1."""
    rows = [f"{40 + index},note number {index} of the visit" for index in range(30)]
    path = tmp_path / "demo.csv"
    path.write_text("Age,remark\n" + "\n".join(rows) + "\n", encoding="utf-8")
    remark = profile([path], Options()).config["datasets"][0]["sheets"][0]["columns"][1]
    assert remark["kind"] == "text" and remark["lengths"] == [1, 20]


def test_workbook_identifiers_of_numbers_stay_numbers(tmp_path: Path) -> None:
    """A workbook ID column of integers keeps number cells and its shapes, in a private table too."""
    workbook = openpyxl.Workbook()
    workbook.active.append(["PatientID", "visit"])
    for index in range(1, 31):
        workbook.active.append([1000 + index, f"visit_{index % 3}"])
    workbook.save(tmp_path / "visits.xlsx")
    workbook = openpyxl.Workbook()
    workbook.active.append(["PatientID", "age"])
    for index in range(1, 31):
        workbook.active.append([index, 20 + index])
    workbook.save(tmp_path / "ages.xlsx")
    result = profile([tmp_path / "visits.xlsx", tmp_path / "ages.xlsx"], Options())
    visits, ages = (dataset["sheets"][0]["columns"][0] for dataset in result.config["datasets"])
    assert visits == {"name": "PatientID", "kind": "identifier", "pattern": "9999",
                      "as_number": True}
    assert ages["pattern"] == {"99": 0.7, "9": 0.3} and ages["as_number"] is True
    config = tmp_path / "config.yaml"
    config.write_text(to_yaml(result), encoding="utf-8")
    generate(config, tmp_path / "out")
    for name in ("file_1.xlsx", "file_2.xlsx"):
        cells = read_tables(tmp_path / "out" / name)[0].column(0)
        assert all(isinstance(cell, int) for cell in cells) and len(set(cells)) == 30


def test_number_identifier_needs_a_pattern_of_digits() -> None:
    """as_number with a pattern of letters is a config error."""
    data = yaml.safe_load(EXAMPLE.read_text(encoding="utf-8"))
    data["datasets"][0]["sheets"][0]["columns"][0]["as_number"] = True
    with pytest.raises(ConfigError, match="digits only"):
        parse_config(data)


def test_percent_column_reads_as_numbers_and_keeps_its_sign(tmp_path: Path) -> None:
    """A column of values such as 85% is numeric with the suffix %, and the synthetic cells keep the sign."""
    rows = [f"P{index:03d},{40 + index % 55}%" for index in range(80)] + ["P999,NA"]
    path = tmp_path / "scores.csv"
    path.write_text("patient_id,score\n" + "\n".join(rows) + "\n", encoding="utf-8")
    report = glimpse([path], Options())
    score = fact_map(report.tables[0].columns[1].facts)
    assert report.tables[0].columns[1].kind == "numeric" and score["suffix"] == "'%'"
    result = profile([path], Options())
    column = result.config["datasets"][0]["sheets"][0]["columns"][1]
    assert column["suffix"] == "%" and column["tokens"] == {"NA": 0.0123}
    config = tmp_path / "config.yaml"
    config.write_text(to_yaml(result), encoding="utf-8")
    generate(config, tmp_path / "out")
    cells = [cell for cell in read_tables(tmp_path / "out" / "file_1.csv")[0].column(1) if cell != "NA"]
    assert all(re.fullmatch(r"\d+%", cell) for cell in cells)
    assert all(40 <= int(cell[:-1]) <= 94 for cell in cells)
    synthetic = glimpse([tmp_path / "out" / "file_1.csv"], Options())
    assert fact_map(synthetic.tables[0].columns[1].facts)["suffix"] == "'%'"


def test_later_date_comes_after_its_earlier_date(tmp_path: Path) -> None:
    """A discharge column comes on or after the admission of its row, from the profile to the synthetic file."""
    start = dt.datetime(2023, 1, 1, 8, 0)
    rows = []
    for index in range(60):
        admit = start + dt.timedelta(days=index * 3, hours=index % 7)
        discharge = admit + dt.timedelta(days=index % 9, hours=2)
        rows.append(f"A{index:03d},{admit:%Y-%m-%d %H:%M:%S},{discharge:%Y-%m-%d %H:%M:%S}")
    path = tmp_path / "stays.csv"
    path.write_text("stay_id,admit_time,discharge_time\n" + "\n".join(rows) + "\n", encoding="utf-8")
    report = glimpse([path], Options())
    assert fact_map(report.tables[0].columns[2].facts)["after"] == "'admit_time'"
    result = profile([path], Options())
    discharge = result.config["datasets"][0]["sheets"][0]["columns"][2]
    assert discharge["after"] == "admit_time" and discharge["interval_days"][0] >= 0
    config = tmp_path / "config.yaml"
    config.write_text(to_yaml(result), encoding="utf-8")
    generate(config, tmp_path / "out")
    table = read_tables(tmp_path / "out" / "file_1.csv")[0]
    pairs = [(dt.datetime.fromisoformat(admit), dt.datetime.fromisoformat(out))
             for admit, out in zip(table.column(1), table.column(2), strict=True)]
    assert all(out >= admit for admit, out in pairs)
    assert max((out - admit).days for admit, out in pairs) <= 9


def test_after_needs_an_earlier_date_column() -> None:
    """An after field that names a missing, a later or a non-date column is a config error."""
    data = yaml.safe_load(EXAMPLE.read_text(encoding="utf-8"))
    columns = data["datasets"][0]["sheets"][0]["columns"]
    columns[4]["after"] = "visit_end"
    with pytest.raises(ConfigError, match="no earlier column is named 'visit_end'"):
        parse_config(data)
    columns[4]["after"] = "feature_01"
    with pytest.raises(ConfigError, match="is not a date column"):
        parse_config(data)


def write_patients_and_admissions(folder: Path) -> list[Path]:
    """Write an invented patient table and an admission table that refers to it.

    Args:
        folder: The output folder.

    Returns:
        The two files.
    """
    patients = folder / "patients.csv"
    patients.write_text("patient_id,arm\n" + "\n".join(f"P{1000 + 7 * index},{'ab'[index % 2]}"
                                                       for index in range(40)) + "\n", encoding="utf-8")
    admissions = folder / "admissions.csv"
    admissions.write_text("admission_id,patient_id\n" + "\n".join(f"V{5000 + index},P{1000 + 7 * (index % 25)}"
                                                                  for index in range(90)) + "\n", encoding="utf-8")
    return [patients, admissions]


def test_linked_identifier_draws_the_keys_of_another_table(tmp_path: Path) -> None:
    """Each synthetic admission names a synthetic patient, as each real one does."""
    files = write_patients_and_admissions(tmp_path)
    report = glimpse(files, Options())
    assert fact_map(report.tables[1].columns[1].facts)["keys of"] == "file_1.csv 'patient_id'"
    result = profile(files, Options())
    linked = result.config["datasets"][1]["sheets"][0]["columns"][1]
    assert linked["link"] == {"dataset": "file_1", "column": "patient_id"} and linked["distinct"] == 25
    config = tmp_path / "config.yaml"
    config.write_text(to_yaml(result), encoding="utf-8")
    generate(config, tmp_path / "out")
    keys = set(read_tables(tmp_path / "out" / "file_1.csv")[0].column(0))
    children = read_tables(tmp_path / "out" / "file_2.csv")[0].column(1)
    assert set(children) <= keys and len(set(children)) == 25


@pytest.mark.parametrize(("link", "message"), [
    ({"dataset": "missing", "column": "patient_id"}, "no CSV, TSV, Excel or Parquet dataset"),
    ({"dataset": "file_1", "column": "arm"}, "no identifier column 'arm'"),
    ({"dataset": "file_2", "column": "admission_id"}, "a link names a column of another table"),
])
def test_broken_link_is_a_config_error(tmp_path: Path, link: dict, message: str) -> None:
    """A link that names a missing dataset, a column that is not an identifier or its own table fails."""
    result = profile(write_patients_and_admissions(tmp_path), Options())
    result.config["datasets"][1]["sheets"][0]["columns"][1]["link"] = link
    with pytest.raises(ConfigError, match=message):
        parse_config(result.config)


def test_iso_dates_with_a_z_keep_their_format_and_order(tmp_path: Path) -> None:
    """Dates such as 2010-01-23T17:45:28Z, as Synthea writes them, read as dates and keep their order."""
    start = dt.datetime(2010, 1, 23, 17, 45, 28)
    rows = [f"E{index:03d},{start + dt.timedelta(days=index):%Y-%m-%dT%H:%M:%SZ},"
            f"{start + dt.timedelta(days=index, minutes=15 + index):%Y-%m-%dT%H:%M:%SZ}" for index in range(40)]
    path = tmp_path / "encounters.csv"
    path.write_text("Id,START,STOP\n" + "\n".join(rows) + "\n", encoding="utf-8")
    result = profile([path], Options())
    stop = result.config["datasets"][0]["sheets"][0]["columns"][2]
    assert stop["kind"] == "date" and stop["date_format"] == "%Y-%m-%dT%H:%M:%SZ" and stop["after"] == "START"
    config = tmp_path / "config.yaml"
    config.write_text(to_yaml(result), encoding="utf-8")
    generate(config, tmp_path / "out")
    table = read_tables(tmp_path / "out" / "file_1.csv")[0]
    assert all(re.fullmatch(r"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", cell) for cell in table.column(1))
    assert all(end >= begin for begin, end in zip(table.column(1), table.column(2), strict=True))


def test_uuid_keys_stay_uuids_and_link(tmp_path: Path) -> None:
    """A column of UUID keys gets random UUIDs, and a linked column draws them."""
    patients = [str(uuid.UUID(int=index * 7919 + 1)) for index in range(30)]
    (tmp_path / "patients.csv").write_text("Id,arm\n" + "\n".join(f"{key},{'ab'[index % 2]}" for index, key
                                                                 in enumerate(patients)) + "\n", encoding="utf-8")
    (tmp_path / "visits.csv").write_text("Id,PATIENT\n" + "\n".join(f"{uuid.UUID(int=10**6 + index)},"
                                                                    f"{patients[index % 20]}" for index in range(60))
                                         + "\n", encoding="utf-8")
    files = [tmp_path / "patients.csv", tmp_path / "visits.csv"]
    report = glimpse(files, Options())
    assert fact_map(report.tables[0].columns[0].facts)["shapes"] == "uuid"
    result = profile(files, Options())
    assert result.config["datasets"][0]["sheets"][0]["columns"][0]["pattern"] == "uuid"
    config = tmp_path / "config.yaml"
    config.write_text(to_yaml(result), encoding="utf-8")
    generate(config, tmp_path / "out")
    keys = read_tables(tmp_path / "out" / "file_1.csv")[0].column(0)
    assert all(uuid.UUID(key).version == 4 for key in keys) and len(set(keys)) == 30
    assert not set(keys) & set(patients)
    assert set(read_tables(tmp_path / "out" / "file_2.csv")[0].column(1)) <= set(keys)
    synthetic = glimpse([tmp_path / "out" / "file_1.csv", tmp_path / "out" / "file_2.csv"], Options())
    assert [item for item in compare(report, synthetic) if item.level == "error"] == []


def test_zip_codes_keep_their_leading_zero(tmp_path: Path) -> None:
    """Digit codes with a leading zero, such as Massachusetts ZIP codes, are codes and keep the zero."""
    rows = [f"P{index:03d},0{1000 + index * 37 % 1700}" for index in range(80)]
    path = tmp_path / "addresses.csv"
    path.write_text("patient_id,ZIP\n" + "\n".join(rows) + "\n", encoding="utf-8")
    report = glimpse([path], Options())
    zip_column = report.tables[0].columns[1]
    assert zip_column.kind == "identifier" and fact_map(zip_column.facts)["shapes"] == "'09999' 80"
    result = profile([path], Options())
    assert result.config["datasets"][0]["sheets"][0]["columns"][1]["pattern"] == "09999"
    config = tmp_path / "config.yaml"
    config.write_text(to_yaml(result), encoding="utf-8")
    generate(config, tmp_path / "out")
    cells = [str(cell) for cell in read_tables(tmp_path / "out" / "file_1.csv")[0].column(1)]
    assert all(re.fullmatch(r"0[1-9]\d{3}", cell) for cell in cells)
