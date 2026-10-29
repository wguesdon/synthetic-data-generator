"""Shared fixtures: an invented delivery with the quirks of real ones.

Every value here is invented. The files hold common quirks of
data files: a semicolon delimiter, a level with a trailing space, NA
tokens in a numeric column, exact zeros, numbers stored as text in a workbook,
and a clinical table that must print counts only.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl
import pytest

from sdg.tables import write_text_table

SUBJECTS = [f"PLT-{index:04d}" for index in range(1, 17)]
CONDITIONS = ["neat", "diluted "]


def plate_rows() -> list[list[str]]:
    """Build the rows of the invented plate export.

    Returns:
        One row per subject and condition, 32 rows.
    """
    rows = []
    for index, (subject, condition) in enumerate((s, c) for s in SUBJECTS for c in CONDITIONS):
        if index % 7 == 3:
            value = "NA"
        elif index % 11 == 5:
            value = "0"
        else:
            value = f"{(index * 37 % 97) / 3.1:.2f}"
        acquired = "TRUE" if index % 4 else "FALSE"
        run_date = (dt.date(2025, 1, 20) + dt.timedelta(days=index % 9)).isoformat()
        rows.append([subject, condition, value, f"reader note {index} for the sample", acquired, run_date])
    return rows


def write_clinical(path: Path) -> None:
    """Write the invented clinical workbook, which must stay private.

    Args:
        path: The workbook path.
    """
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Demographics"
    sheet.append(["Subject", "Age", "Sex", "Cohort"])
    for index, subject in enumerate(SUBJECTS):
        sheet.append([subject, 30 + index, "F" if index % 2 else "M", ["group_1", "group_2"][index % 2]])
    workbook.save(path)


def write_assay(path: Path) -> None:
    """Write the invented assay workbook with numbers stored as text.

    Args:
        path: The workbook path.
    """
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    sheet.append(["Plate", "IL6", "Measured on"])
    for index in range(24):
        sheet.append([f"P0{index % 2 + 1}", f"{1.5 + index * 0.25:.3f}",
                      dt.datetime(2026, 9, 1) + dt.timedelta(days=index)])
    workbook.save(path)


@pytest.fixture
def delivery(tmp_path: Path) -> Path:
    """Write the invented delivery into a folder.

    Args:
        tmp_path: The pytest temporary folder.

    Returns:
        The delivery folder, with ``plate.csv``, ``clinical.xlsx`` and ``assay.xlsx``.
    """
    folder = tmp_path / "delivery"
    folder.mkdir()
    write_text_table(folder / "plate.csv", ["Sample ID", "Dilution", "OD450", "Comment", "Detected",
                                            "Run date"], plate_rows(), ";")
    write_clinical(folder / "clinical.xlsx")
    write_assay(folder / "assay.xlsx")
    return folder
