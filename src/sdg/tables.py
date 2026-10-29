"""Read and write delivered tables cell by cell, and classify their columns.

Tables are read with the ``csv`` module and ``openpyxl`` rather than pandas.
pandas renames a repeated column name (``a`` becomes ``a.1``), converts types
and drops blank lines, and a synthetic file must copy exactly those quirks.
Every cell comes back as delivered: text from a CSV file, and the stored Python
type (str, int, float, bool, datetime) from a workbook or a Parquet file.
"""

from __future__ import annotations

import csv
import datetime as dt
import math
import re
import statistics
from collections import Counter
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl
import pandas as pd

from sdg import fake_values, privacy
from sdg.report import ColumnReport, Fact, counted

#: Candidate delimiters of a text table, in order of preference on a tie.
DELIMITERS = (",", ";", "\t", "|")

#: Lines of a text table that choose its delimiter, comment lines excluded.
DELIMITER_LINES = 50

#: A first cell that starts with one of these marks a comment row. cBioPortal
#: and featureCounts files use ``#``; a GEO series matrix uses ``!``.
COMMENT_MARKS = ("#", "!")

#: Rows searched for the header row, after the leading comment rows.
HEADER_SCAN = 50

#: A header row fills more than this share of the table width. A title row,
#: such as a study name in cell A1, fills less.
HEADER_FILL = 0.5

#: Rows that set the table width for the header search.
WIDTH_ROWS = 1000

#: Rows below the header candidates that give the kind of each column.
DATA_SAMPLE = 100

#: Column kinds that confirm where the data starts. The synthetic file keeps
#: these kinds, so it reads with the same header row.
TYPED_KINDS = {"numeric", "date", "boolean"}

#: A cell with a digit can be a date; a cell without one is not tried.
DIGIT = re.compile(r"\d")

#: Number of quantiles that a profile records: the deciles.
QUANTILE_POINTS = 11

#: Text date formats tried on a column, in order.
DATE_FORMATS = ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M:%SZ", "%d/%m/%Y", "%m/%d/%Y",
                "%d-%m-%Y", "%Y/%m/%d", "%d.%m.%Y", "%d-%b-%Y", "%d%b%Y", "%Y%m%d")

#: Date formats with no separator, which also read a numeric identifier.
COMPACT_DATE_FORMATS = {"%Y%m%d", "%d%b%Y"}

#: A column is a date column when at least this share of its filled cells are dates.
DATE_SHARE = 0.9

#: Text cells read as a boolean.
BOOLEAN_TEXT = {"true", "false"}

#: Most shapes that a report prints for a text column, and that a profile keeps.
TEXT_SHAPES = 8

#: A text column keeps its common shapes when they fill at least this share of its cells.
TEXT_SHAPE_SHARE = 0.5

#: A number with a percent sign, such as ``85%`` or ``12.5 %``. The second
#: group is the suffix, with its space.
PERCENT = re.compile(r"^\s*[-+]?(\d+\.?\d*|\.\d+)(\s*%)\s*$")

#: Most decimal places recorded for a numeric column.
MAX_DECIMALS = 6


@dataclass
class Table:
    """One table as delivered.

    Attributes:
        sheet: Sheet name for a workbook, ``None`` for a text file.
        header: Column names in file order, repeats included.
        rows: Data rows, each padded or cut to the header length.
        delimiter: Delimiter of a text file, ``None`` for a workbook.
        ragged: Count of data rows whose cell count differed from the header.
        preamble: Rows above the header, such as a title, comment lines or
            a second header row, as read.
        footer: Comment rows below the data, such as the end line of a GEO
            series matrix, as read.
    """

    sheet: str | None
    header: list[str]
    rows: list[list[object]]
    delimiter: str | None = None
    ragged: int = 0
    preamble: list[list[object]] = field(default_factory=list)
    footer: list[list[object]] = field(default_factory=list)

    def column(self, index: int) -> list[object]:
        """Return the cells of one column.

        Args:
            index: Column position.

        Returns:
            The cells, one per data row.
        """
        return [row[index] for row in self.rows]


def read_text(path: Path) -> str:
    """Read a text file, with a fallback for files that are not UTF-8.

    Args:
        path: The file.

    Returns:
        The text, without a byte order mark.
    """
    data = path.read_bytes()
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def sniff_delimiter(first_line: str, suffix: str) -> str:
    """Choose the delimiter of a text table from its header line.

    The header rarely holds a decimal comma, so the count of each candidate in
    it is a safer guide than ``csv.Sniffer`` on the data rows.

    Args:
        first_line: The header line.
        suffix: File suffix. ``.tsv`` always means a tab.

    Returns:
        The delimiter.
    """
    if suffix == ".tsv":
        return "\t"
    counts = {delimiter: first_line.count(delimiter) for delimiter in DELIMITERS}
    best = max(DELIMITERS, key=counts.__getitem__)
    return best if counts[best] else ","


def is_comment_line(line: str) -> bool:
    """Say whether a line of a text file is a comment line.

    Args:
        line: One line, as read.

    Returns:
        True when the line starts with a mark of :data:`COMMENT_MARKS`.
    """
    return line.startswith(COMMENT_MARKS)


def table_delimiter(lines: list[str], suffix: str) -> str:
    """Choose the delimiter of a text table from its first lines.

    Blank lines and comment lines do not count. The delimiter that splits the
    most lines into the same number of cells wins, because a title line, a
    decimal comma or a featureCounts cell such as ``chr1;chr1`` changes the
    count of some lines only. The first line of a featureCounts file is a
    comment that holds a semicolon and no tab.

    Args:
        lines: The lines of the file.
        suffix: File suffix. ``.tsv`` always means a tab.

    Returns:
        The delimiter. A tie goes to the earlier delimiter of
        :data:`DELIMITERS`, and a file with no candidate gets a comma.
    """
    if suffix == ".tsv":
        return "\t"
    sample: list[str] = []
    for line in lines:
        if line.strip() and not is_comment_line(line):
            sample.append(line)
            if len(sample) == DELIMITER_LINES:
                break
    best, best_lines = ",", 0
    for delimiter in DELIMITERS:
        counts = Counter(line.count(delimiter) for line in sample)
        counts.pop(0, None)
        shared = max(counts.values(), default=0)
        if shared > best_lines:
            best, best_lines = delimiter, shared
    return best


def is_comment_row(row: list[object]) -> bool:
    """Say whether a row is a comment row.

    Args:
        row: A row as read.

    Returns:
        True when the first cell is text that starts with a mark of
        :data:`COMMENT_MARKS` and is not a missing token: a TCGA-CDR column
        can end with many ``#N/A`` cells.
    """
    return (bool(row) and isinstance(row[0], str) and row[0].startswith(COMMENT_MARKS)
            and not privacy.is_missing_token(row[0]))


def is_blank(value: object) -> bool:
    """Say whether a cell holds nothing to read: empty, or spaces only.

    Args:
        value: A cell as read.

    Returns:
        True for an empty cell or a cell of spaces.
    """
    return is_empty(value) or (isinstance(value, str) and not value.strip())


def is_label(value: object) -> bool:
    """Say whether a filled cell can be a column name.

    A name is text that is not a number, a date, a boolean or a standard
    missing token. A data row of text only, such as a row of levels, is all
    labels too, so a row of labels is only a header candidate.

    Args:
        value: A filled cell.

    Returns:
        True for a label.
    """
    if not isinstance(value, str):
        return False
    text = value.strip()
    if (number(text) is not None or text.lower() in BOOLEAN_TEXT or privacy.is_missing_token(text)):
        return False
    if DIGIT.search(text):
        return all(parse_date(text, date_format) is None for date_format in DATE_FORMATS)
    return True


def filled_count(row: list[object]) -> int:
    """Count the filled cells of a row.

    Args:
        row: A row as read.

    Returns:
        The count of cells that are not empty.
    """
    return sum(not is_empty(cell) for cell in row)


def row_span(row: list[object]) -> int:
    """Return the width that a row fills: the position after its last filled cell.

    Args:
        row: A row as read.

    Returns:
        The span, 0 for an empty row.
    """
    for position in range(len(row) - 1, -1, -1):
        if not is_empty(row[position]):
            return position + 1
    return 0


def is_label_row(row: list[object]) -> bool:
    """Say whether a row holds labels only.

    Args:
        row: A row as read.

    Returns:
        True when the row has a label and no other value.
    """
    cells = [cell for cell in row if not is_blank(cell)]
    return bool(cells) and all(is_label(cell) for cell in cells)


def typed_columns(records: list[list[object]], first: int) -> list[int]:
    """Find the columns of numbers, dates or booleans from a row down.

    The kind of a whole column decides, as in a glimpse, and not one cell. The
    TCGA-CDR TSS sheet has a text column of codes such as ``01``, and the
    synthetic file writes invented words there, so a header row that only
    such a cell confirms would not read the same in the synthetic file.

    Args:
        records: The rows.
        first: Position of the first row of the sample.

    Returns:
        The positions of the typed columns.
    """
    sample = records[first:first + DATA_SAMPLE]
    width = max((len(row) for row in sample), default=0)
    typed = []
    for position in range(width):
        cells = [row[position] if position < len(row) else None for row in sample]
        if analyze_column("", cells).kind in TYPED_KINDS:
            typed.append(position)
    return typed


def find_header(records: list[list[object]]) -> int:
    """Find the header row of a table.

    The leading comment rows and blank rows go first. Then a row that fills
    at most :data:`HEADER_FILL` of the width, such as a title, is skipped. The
    candidates are the next rows of labels that fill the width; a blank row
    between them does not end them. The rows after them must hold a column of
    numbers, dates or booleans, which confirms where the data starts. The
    header is then the candidate with the most filled cells that names each
    such column: the UCI credit workbook has a row of codes ``X1`` to ``X23``,
    with cell A1 empty, above the row of names. With no such column, as in a
    table of text only, the header is the first row after the comment rows.
    With no candidate, the header is the first row.

    Args:
        records: The rows, trailing empty rows removed.

    Returns:
        The position of the header row.
    """
    start = 0
    while start < len(records) and (is_comment_row(records[start]) or not filled_count(records[start])):
        start += 1
    width = max((row_span(row) for row in records[start:start + WIDTH_ROWS]), default=0)
    candidates: list[int] = []
    after = None
    for index in range(start, min(len(records), start + HEADER_SCAN)):
        row = records[index]
        filled = filled_count(row)
        if not filled or (not candidates and is_comment_row(row)):
            continue
        wide = filled > HEADER_FILL * width
        if wide and is_label_row(row):
            candidates.append(index)
        elif candidates:
            after = index
            break
        elif wide:
            # A wide row of values before any label row: the table has no
            # header row to find, so the first row stays the header.
            break
    if not candidates:
        return 0
    typed = typed_columns(records, after) if after is not None else []
    if not typed:
        return start
    for _, index in sorted((-filled_count(records[index]), index) for index in candidates):
        row = records[index]
        if all(position < len(row) and not is_empty(row[position]) for position in typed):
            return index
    return candidates[0]


def data_end(records: list[list[object]], start: int) -> int:
    """Find the end of the data rows, before the comment rows of a footer.

    A GEO series matrix ends with the line ``!series_matrix_table_end``. The
    trailing comment rows, and the blank rows above them, are the footer. A
    column of values such as ``#12`` makes each data row a comment row, so
    a footer needs data rows that are not comment rows.

    Args:
        records: The rows, trailing empty rows removed.
        start: Position of the first data row.

    Returns:
        The position after the last data row.
    """
    end = len(records)
    while end > start and is_comment_row(records[end - 1]):
        end -= 1
    if end in (len(records), start) or any(is_comment_row(row) for row in records[start:end]):
        return len(records)
    while end > start and not filled_count(records[end - 1]):
        end -= 1
    return end


def strip_trailing_empty(rows: list[list[object]]) -> list[list[object]]:
    """Remove the fully empty rows at the end of a table.

    A spreadsheet often reports formatted but empty rows after the data, and a
    text file often ends with a blank line. Empty rows inside the data stay.

    Args:
        rows: Rows as read.

    Returns:
        The rows without the trailing empty ones.
    """
    end = len(rows)
    while end and all(is_empty(cell) for cell in rows[end - 1]):
        end -= 1
    return rows[:end]


def fit_rows(header: list[str], rows: list[list[object]]) -> tuple[list[list[object]], int]:
    """Pad or cut each row to the header length.

    Args:
        header: Column names.
        rows: Data rows as read.

    Returns:
        The fitted rows and the count of rows that needed a change.
    """
    width = len(header)
    fitted, ragged = [], 0
    for row in rows:
        if len(row) != width and any(not is_empty(cell) for cell in row):
            ragged += 1
        fitted.append((list(row) + [""] * width)[:width])
    return fitted, ragged


def read_text_table(path: Path, delimiter: str | None = None) -> Table:
    """Read a CSV or TSV file as delivered text.

    Args:
        path: The file.
        delimiter: Delimiter to use. ``None`` chooses it from the first lines.

    Returns:
        The table, with the rows above its header and its footer apart.
    """
    lines = read_text(path).splitlines()
    delimiter = delimiter or table_delimiter(lines, privacy.file_suffix(path))
    records = strip_trailing_empty(list(csv.reader(lines, delimiter=delimiter)))
    if not records:
        return Table(None, [], [], delimiter)
    top = find_header(records)
    end = data_end(records, top + 1)
    header = [str(name) for name in records[top]]
    rows, ragged = fit_rows(header, records[top + 1:end])
    return Table(None, header, rows, delimiter, ragged, records[:top], records[end:])


def read_workbook(path: Path) -> list[Table]:
    """Read every sheet of an Excel workbook with the stored cell types.

    Formula cells give their cached value. Trailing empty rows go, and so do
    trailing columns of the table with no name and no value.

    Args:
        path: The ``.xlsx`` file.

    Returns:
        One table per sheet, in workbook order, with the rows above each
        header and each footer apart.

    Raises:
        ValueError: If the file is an old ``.xls`` workbook.
    """
    if privacy.file_suffix(path) == ".xls":
        raise ValueError(f"{path.name}: save the .xls workbook as .xlsx, which openpyxl reads")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    tables = []
    try:
        for sheet in workbook.worksheets:
            records = strip_trailing_empty(sheet_records(sheet))
            if not records:
                tables.append(Table(sheet.title, [], []))
                continue
            top = find_header(records)
            end = data_end(records, top + 1)
            body = records[top:end]
            width = max(len(record) for record in body)
            while width and all(is_empty(record[width - 1]) if len(record) >= width else True
                                for record in body):
                width -= 1
            header = ["" if cell is None else str(cell) for cell in body[0][:width]]
            rows, ragged = fit_rows(header, [record[:width] for record in body[1:]])
            tables.append(Table(sheet.title, header, [[None if cell == "" else cell for cell in row]
                                                      for row in rows], None, ragged,
                                records[:top], records[end:]))
    finally:
        workbook.close()
    return tables


def sheet_records(sheet: Any) -> list[list[object]]:
    """Read every row of a read-only worksheet, padded to one width.

    The read-only mode trusts the dimension tag of the sheet, and some writers
    store a wrong tag such as ``A1:A1``, which hides every other cell.
    Resetting the dimensions reads each cell that the sheet holds.

    Args:
        sheet: A worksheet of a workbook opened with ``read_only=True``.

    Returns:
        The rows, each as long as the longest row.
    """
    sheet.reset_dimensions()
    records = [list(row) for row in sheet.iter_rows(values_only=True)]
    width = max((len(record) for record in records), default=0)
    return [record + [None] * (width - len(record)) for record in records]


def parquet_cell(value: object) -> object:
    """Turn one Parquet cell into a cell of the types that a workbook gives.

    Args:
        value: The cell from ``to_pylist``.

    Returns:
        The cell: None, a bool, an int, a float, a text, a date or a naive
        datetime. A decimal becomes a float, a time zone becomes UTC, and a
        nested or binary value becomes text.
    """
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    if hasattr(value, "to_pydatetime"):
        value = value.to_pydatetime()
    if isinstance(value, dt.datetime):
        return value.astimezone(dt.UTC).replace(tzinfo=None) if value.tzinfo else value
    if isinstance(value, dt.date):
        return value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    return str(value)


def read_parquet(path: Path) -> Table:
    """Read a Parquet file as one table with typed cells.

    A Parquet file has no rows above the header and no footer, and a null is
    an empty cell.

    Args:
        path: The ``.parquet`` file.

    Returns:
        The table.
    """
    import pyarrow.parquet as parquet

    data = parquet.read_table(path)
    columns = [[parquet_cell(value) for value in column.to_pylist()] for column in data.columns]
    rows = [list(row) for row in zip(*columns, strict=True)] if columns else []
    return Table(None, [str(name) for name in data.column_names], rows)


def write_parquet(path: Path, header: list[str], arrays: list[Any]) -> None:
    """Write one table as a Parquet file.

    Args:
        path: The output file.
        header: Column names in order.
        arrays: One ``pyarrow`` array per column.
    """
    import pyarrow as pa
    import pyarrow.parquet as parquet

    parquet.write_table(pa.Table.from_arrays(arrays, names=header), path)


def frame_cell(value: object) -> object:
    """Turn one DataFrame cell into a cell as the file readers give it.

    Args:
        value: The cell.

    Returns:
        None for a missing value; a Python bool, int, float or datetime for a
        NumPy or pandas scalar; the value itself otherwise.
    """
    if value is None or (not isinstance(value, (str, bytes, list, tuple, dict)) and pd.isna(value)):
        return None
    if isinstance(value, (bool, np.bool_)):
        return bool(value)
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        return float(value)
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, (int, float, str, dt.date)):
        return value
    return str(value)


def frame_table(frame: pd.DataFrame) -> Table:
    """Turn a DataFrame into a table.

    An index that is not the default range holds data, such as sample names,
    so it becomes the first column.

    Args:
        frame: The DataFrame.

    Returns:
        The table.
    """
    if not isinstance(frame.index, pd.RangeIndex):
        frame = frame.reset_index()
    rows = [[frame_cell(value) for value in row] for row in frame.itertuples(index=False, name=None)]
    return Table(None, [str(name) for name in frame.columns], rows)


#: Text of the error when the extra of the RDS format is not installed.
MISSING_R_EXTRA = 'the rds format needs pyreadr; run uv add "sdg[r]"'


def import_pyreadr() -> Any:
    """Import pyreadr, the optional dependency of the RDS format.

    Returns:
        The module.

    Raises:
        ValueError: If the extra is not installed.
    """
    try:
        import pyreadr
    except ImportError as error:
        raise ValueError(MISSING_R_EXTRA) from error
    return pyreadr


def read_rds(path: Path) -> Table:
    """Read an RDS file of one R data frame as one table with typed cells.

    Args:
        path: The ``.rds`` file.

    Returns:
        The table. A factor reads as text, and an R date as a date.
    """
    frames = import_pyreadr().read_r(str(path))
    return frame_table(next(iter(frames.values())).reset_index(drop=True))


def write_rds(path: Path, header: list[str], columns: list[list[object]]) -> None:
    """Write one table as an RDS file of an R data frame.

    Args:
        path: The output file.
        header: Column names in order.
        columns: The cells of each column: None, a bool, a number or a text.
    """
    frame = pd.DataFrame({name: pd.Series(cells, dtype=object) for name, cells in zip(header, columns, strict=True)})
    for name in frame.columns:
        values = [cell for cell in frame[name] if cell is not None]
        if values and all(isinstance(cell, bool) for cell in values):
            frame[name] = frame[name].astype("boolean")
        elif values and all(isinstance(cell, (int, float)) and not isinstance(cell, bool) for cell in values):
            frame[name] = pd.to_numeric(frame[name])
    import_pyreadr().write_rds(str(path), frame)


def read_tables(path: Path, delimiter: str | None = None) -> list[Table]:
    """Read every table of a CSV, TSV, Excel or Parquet file.

    Args:
        path: The file.
        delimiter: Delimiter of a text file. ``None`` sniffs it.

    Returns:
        The tables: one for a text file or a Parquet file, one per sheet for a
        workbook.
    """
    suffix = privacy.file_suffix(path)
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        return read_workbook(path)
    if suffix == ".parquet":
        return [read_parquet(path)]
    if suffix == ".rds":
        return [read_rds(path)]
    return [read_text_table(path, delimiter)]


def is_empty(value: object) -> bool:
    """Say whether a cell is empty.

    Args:
        value: A cell as read.

    Returns:
        True for a missing cell or an empty string. A cell of spaces is not
        empty, because a stray space is a delivered quirk.
    """
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    return isinstance(value, str) and value == ""


def number(value: object) -> float | None:
    """Parse one delivered cell as a finite number.

    Args:
        value: A cell as read.

    Returns:
        The number, or None when the cell is not a finite number. The texts
        ``nan`` and ``inf`` are not numbers here: they are tokens.
    """
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value) if math.isfinite(value) else None
    if not isinstance(value, str):
        return None
    try:
        parsed = float(value.strip())
    except ValueError:
        return None
    return parsed if math.isfinite(parsed) else None


def decimals_of(value: object) -> int | None:
    """Count the decimal places of a number as delivered.

    Args:
        value: A numeric cell.

    Returns:
        The decimal places, or None for a number in exponent form.
    """
    text = value.strip() if isinstance(value, str) else repr(value)
    if "e" in text.lower():
        return None
    if isinstance(value, float) and value.is_integer() and not isinstance(value, str):
        return 0
    return len(text.split(".", 1)[1]) if "." in text else 0


def parse_date(value: object, date_format: str | None) -> dt.datetime | None:
    """Parse one cell as a date.

    Args:
        value: A cell as read.
        date_format: ``strptime`` format for a text cell. A workbook date cell
            needs none.

    Returns:
        The date as a datetime, or None.
    """
    if isinstance(value, dt.datetime):
        return value
    if isinstance(value, dt.date):
        return dt.datetime(value.year, value.month, value.day)
    if not isinstance(value, str) or date_format is None:
        return None
    try:
        return dt.datetime.strptime(value.strip(), date_format)
    except ValueError:
        return None


def find_date_format(filled: list[object]) -> tuple[bool, str | None]:
    """Decide whether a column holds dates, and in which text format.

    Args:
        filled: The filled cells of the column.

    Returns:
        Whether the column is a date column, and the text format. The format is
        None for workbook date cells.
    """
    if not filled:
        return False, None
    native = sum(isinstance(value, (dt.date, dt.datetime)) for value in filled)
    if native >= DATE_SHARE * len(filled):
        return True, None
    texts = [value for value in filled if isinstance(value, str)]
    if len(texts) < DATE_SHARE * len(filled):
        return False, None
    for date_format in DATE_FORMATS:
        parsed = sum(parse_date(value, date_format) is not None for value in texts)
        if parsed >= DATE_SHARE * len(filled):
            return True, date_format
    return False, None


@dataclass
class ColumnStats:
    """What one column holds, before any privacy rule applies.

    This object never leaves the secure machine. A glimpse report or a config
    takes only the parts that the privacy rules allow.

    Attributes:
        name: Column name.
        kind: ``identifier``, ``numeric``, ``categorical``, ``text``, ``date``
            or ``boolean``.
        filled: Count of filled cells.
        empty: Count of empty cells.
        unique: Count of distinct filled values.
        cell_types: Python type names of the filled cells and their counts.
        numbers: The numeric cells.
        tokens: Non-numeric filled cells of a numeric column and their counts.
        values: Text of each filled cell and its count.
        date_format: Text date format, or None.
        dates: The parsed dates.
        decimals: Most decimal places of a number, or None.
        suffix: Text after each number, such as ``%``, or None.
        as_text: Numbers or dates stored as text in a workbook.
        true_count: Count of true cells of a boolean column.
    """

    name: str
    kind: str
    filled: int
    empty: int
    unique: int
    cell_types: Counter
    numbers: list[float] = field(default_factory=list)
    tokens: Counter = field(default_factory=Counter)
    values: Counter = field(default_factory=Counter)
    date_format: str | None = None
    dates: list[dt.datetime] = field(default_factory=list)
    decimals: int | None = None
    suffix: str | None = None
    as_text: bool = False
    true_count: int = 0

    @property
    def rows(self) -> int:
        """Return the count of cells in the column.

        Returns:
            Filled plus empty cells.
        """
        return self.filled + self.empty


def is_boolean(value: object) -> bool:
    """Say whether a cell is a boolean or the text TRUE or FALSE.

    Args:
        value: A filled cell.

    Returns:
        True for a boolean cell.
    """
    return isinstance(value, bool) or (isinstance(value, str) and value.strip().lower() in BOOLEAN_TEXT)


def is_true(value: object) -> bool:
    """Say whether a boolean cell is true.

    Args:
        value: A boolean cell.

    Returns:
        True for a true cell.
    """
    return value is True or (isinstance(value, str) and value.strip().lower() == "true")


def column_kind(name: str, filled: list[object], numbers: list[float], tokens: Counter,
                values: Counter, max_categories: int) -> tuple[str, str | None]:
    """Choose the kind of a column from its cells.

    Args:
        name: Column name.
        filled: Filled cells.
        numbers: The numeric cells.
        tokens: Non-numeric filled cells and their counts.
        values: Text of each filled cell and its count.
        max_categories: Most levels of a categorical column.

    Returns:
        The kind, and the text date format of a date column.
    """
    if not filled:
        return "text", None
    is_date, date_format = find_date_format(filled)
    if privacy.is_identifier_name(name):
        # A "Sample date" column holds dates, but a compact date format such
        # as 20261002 also reads a numeric barcode, so it does not count here.
        if is_date and date_format not in COMPACT_DATE_FORMATS:
            return "date", date_format
        return "identifier", None
    if all(is_boolean(value) for value in filled):
        return "boolean", None
    # A ZIP code such as 01013 is a code: a number would lose its zero.
    codes = sum(isinstance(value, str) and bool(privacy.LEADING_ZERO.fullmatch(value.strip())) for value in filled)
    if numbers and codes >= privacy.CODE_SHARE * len(numbers):
        return "identifier", None
    # The share leaves out standard missing tokens: a TCGA column of days with
    # 68% "#N/A" cells is numeric, not free text.
    missing = sum(count for token, count in tokens.items() if privacy.is_missing_token(token))
    if (numbers and len(numbers) >= privacy.NUMERIC_SHARE * (len(filled) - missing)
            and len(tokens) <= privacy.NUMERIC_TOKENS):
        return "numeric", None
    if is_date:
        return "date", date_format
    if (len(values) <= max_categories and len(values) <= privacy.CATEGORY_DISTINCT_SHARE * len(filled)
            and len(numbers) < privacy.NUMERIC_SHARE * len(filled)):
        return "categorical", None
    return "text", None


def number_suffix(filled: list[object]) -> str | None:
    """Find the percent suffix of a column of numbers such as ``85%``.

    Args:
        filled: The filled cells.

    Returns:
        The most common suffix, ``%`` or `` %``, when at least half of the
        cells are numbers with it, else None.
    """
    found = Counter(match.group(2) for value in filled
                    if isinstance(value, str) and (match := PERCENT.match(value)))
    if not found or sum(found.values()) < privacy.NUMERIC_SHARE * len(filled):
        return None
    return found.most_common(1)[0][0]


def without_suffix(value: object) -> object:
    """Remove the percent suffix of a cell.

    Args:
        value: A filled cell.

    Returns:
        The number text without its suffix, or the cell itself.
    """
    if isinstance(value, str) and PERCENT.match(value):
        return value.strip()[:-1].strip()
    return value


def analyze_column(name: str, cells: list[object], max_categories: int = privacy.MAX_LEVELS) -> ColumnStats:
    """Describe one column fully, for the privacy rules to filter.

    A column of numbers with a percent sign, such as ``85%``, is numeric with
    the suffix ``%``, so its synthetic copy draws numbers and writes the sign.

    Args:
        name: Column name.
        cells: The cells of the column, one per data row.
        max_categories: Most levels of a categorical column.

    Returns:
        The column statistics.
    """
    filled = [value for value in cells if not is_empty(value)]
    suffix = number_suffix(filled)
    if suffix:
        stats = column_stats(name, cells, filled, [without_suffix(value) for value in filled], max_categories)
        if stats.kind == "numeric":
            stats.suffix = suffix
            return stats
    return column_stats(name, cells, filled, filled, max_categories)


def column_stats(name: str, cells: list[object], filled: list[object], readable: list[object],
                 max_categories: int) -> ColumnStats:
    """Describe one column from its filled cells and the cells that read as numbers.

    Args:
        name: Column name.
        cells: The cells of the column, one per data row.
        filled: The filled cells as delivered.
        readable: The filled cells, with a number suffix removed.
        max_categories: Most levels of a categorical column.

    Returns:
        The column statistics.
    """
    parsed = [number(value) for value in readable]
    numbers = [value for value in parsed if value is not None]
    tokens = Counter(str(value) for value, parsed_value in zip(filled, parsed, strict=True)
                     if parsed_value is None)
    values = Counter(str(value) for value in filled)
    kind, date_format = column_kind(name, filled, numbers, tokens, values, max_categories)
    stats = ColumnStats(name=name, kind=kind, filled=len(filled), empty=len(cells) - len(filled),
                        unique=len(values), cell_types=Counter(type(value).__name__ for value in filled),
                        values=values)
    if kind == "numeric":
        stats.numbers = numbers
        stats.tokens = tokens
        places = [decimals_of(value) for value, parsed_value in zip(readable, parsed, strict=True)
                  if parsed_value is not None]
        known = [place for place in places if place is not None]
        stats.decimals = min(max(known), MAX_DECIMALS) if known else None
        numeric_cells = [value for value, parsed_value in zip(filled, parsed, strict=True)
                         if parsed_value is not None]
        stats.as_text = sum(isinstance(value, str) for value in numeric_cells) > len(numeric_cells) / 2
    elif kind == "date":
        stats.date_format = date_format
        stats.dates = [date for value in filled if (date := parse_date(value, date_format)) is not None]
        stats.as_text = date_format is not None
    elif kind == "boolean":
        stats.true_count = sum(is_true(value) for value in filled)
    return stats


def rounded_quantiles(values: list[float]) -> list[float]:
    """Return the deciles of the values, rounded to the report precision.

    Args:
        values: The numbers.

    Returns:
        Eleven rounded quantiles from the minimum to the maximum.
    """
    points = np.quantile(np.asarray(values, dtype=float), np.linspace(0.0, 1.0, QUANTILE_POINTS))
    return [privacy.round_sig(float(point)) for point in points]


def trimmed_quantiles(values: list[float]) -> list[float]:
    """Return eleven quantiles from 5% to 95%, rounded to the report precision.

    A private table records these quantiles instead of the deciles, because
    the minimum and the maximum are each the value of one person.

    Args:
        values: The numbers.

    Returns:
        The rounded quantiles at evenly spaced probabilities from ``TRIM_LOW``
        to ``TRIM_HIGH``.
    """
    points = np.quantile(np.asarray(values, dtype=float),
                         np.linspace(privacy.TRIM_LOW, privacy.TRIM_HIGH, QUANTILE_POINTS))
    return [privacy.round_sig(float(point)) for point in points]


def kept_codes(stats: ColumnStats) -> Counter | None:
    """Return the codes that a numeric column of a private table keeps.

    Args:
        stats: The statistics of a numeric column.

    Returns:
        Each code with ``MIN_VALUES`` cells or more and its count, or None when
        the column has more than ``CODE_LIMIT`` distinct numbers.
    """
    counts = Counter(stats.numbers)
    if not counts or len(counts) > privacy.CODE_LIMIT:
        return None
    return Counter({code: count for code, count in counts.items() if count >= privacy.MIN_VALUES})


def safe_number_facts(stats: ColumnStats, limit: int) -> list[Fact]:
    """Report the safe facts of a numeric column of a private table.

    Args:
        stats: The column statistics.
        limit: Most codes or tokens to list.

    Returns:
        The codes, or the quantiles from 5% to 95% and the median, then the
        decimals and the missing tokens. No minimum and no maximum.
    """
    facts = []
    codes = kept_codes(stats)
    if codes is not None:
        shown: Counter = Counter()
        for code, count in codes.items():
            shown[privacy.format_number(code)] += count
        facts.append(Fact("codes", counted(shown, limit) if shown else "none with 10 cells"))
    elif len(stats.numbers) >= privacy.TRIM_MIN_VALUES:
        low, high = np.quantile(np.asarray(stats.numbers, dtype=float), [privacy.TRIM_LOW, privacy.TRIM_HIGH])
        facts += [Fact("5% to 95%", f"{privacy.format_number(float(low))} to {privacy.format_number(float(high))}"),
                  Fact("median", privacy.format_number(statistics.median(stats.numbers)))]
    else:
        facts.append(Fact("numbers", str(len(stats.numbers))))
    if stats.decimals is not None:
        facts.append(Fact("decimals", str(stats.decimals)))
    if stats.suffix:
        facts.append(Fact("suffix", repr(stats.suffix)))
    missing = Counter({token: count for token, count in stats.tokens.items() if privacy.is_missing_token(token)})
    if missing:
        facts.append(Fact("tokens", counted(missing, limit)))
    facts.append(Fact("non-numeric", str(sum(stats.tokens.values()))))
    return facts


def common_shapes(stats: ColumnStats) -> list[tuple[str, int]]:
    """Return the common shapes of a text column, when they fill half of its cells.

    Args:
        stats: The statistics of a text column.

    Returns:
        The ``TEXT_SHAPES`` most common shapes and their counts, or nothing.
    """
    shapes = Counter(privacy.shape(value) for value in stats.values.elements()).most_common(TEXT_SHAPES)
    return shapes if stats.filled and sum(count for _, count in shapes) >= TEXT_SHAPE_SHARE * stats.filled else []


def text_provider(stats: ColumnStats) -> str | None:
    """Choose the provider of a text column, which its shapes can replace.

    The real shapes give the real form of a code, such as a passport number,
    so a code provider gives way to them.

    Args:
        stats: The statistics of a text column.

    Returns:
        The provider, or None.
    """
    provider = fake_values.column_provider(stats.name, stats.values)
    if provider in fake_values.CODE_PROVIDERS and common_shapes(stats):
        return None
    return provider


def identifier_shapes(stats: ColumnStats, limit: int) -> Fact:
    """Report the shapes of an identifier column.

    A column of UUIDs prints ``uuid``, because the shapes of random hex keys
    are all different and tell nothing.

    Args:
        stats: The statistics of an identifier column.
        limit: Most shapes to list.

    Returns:
        The ``shapes`` fact.
    """
    if stats.values and all(fake_values.UUID.fullmatch(value.strip()) for value in stats.values):
        return Fact("shapes", "uuid")
    return Fact("shapes", counted(Counter(privacy.code_shape(value) for value in stats.values.elements()), limit))


def safe_facts(stats: ColumnStats, limit: int, levels: dict[str, str] | None) -> list[Fact]:
    """Report the safe facts of a column of a private table that is not a private column.

    Args:
        stats: The column statistics.
        limit: Most shapes, codes or levels to list.
        levels: Generic name of each category level.

    Returns:
        Identifier shapes, the safe number facts, the true and false counts
        when each has ``MIN_VALUES`` cells, the date format, the generic level
        counts, or the common shapes of a text column when they fill half of
        its cells. A shape holds no value, as an identifier shape does.
    """
    if stats.kind == "identifier":
        return [identifier_shapes(stats, limit)]
    if stats.kind == "numeric":
        return safe_number_facts(stats, limit)
    if stats.kind == "boolean" and min(stats.true_count, stats.filled - stats.true_count) >= privacy.MIN_VALUES:
        return [Fact("true", str(stats.true_count)), Fact("false", str(stats.filled - stats.true_count))]
    if stats.kind == "date":
        return [Fact("format", repr(stats.date_format) if stats.date_format else "workbook date")]
    if stats.kind == "categorical" and levels is not None:
        return [Fact("values", counted(Counter({levels.get(value, value): count
                                                for value, count in stats.values.items()}), limit))]
    shapes = common_shapes(stats) if stats.kind == "text" else []
    return [Fact("shapes", counted(Counter(dict(shapes)), TEXT_SHAPES))] if shapes else []


def number_facts(stats: ColumnStats, limit: int) -> list[Fact]:
    """Report the facts of a numeric column that the privacy rules allow.

    Args:
        stats: The column statistics.
        limit: Most tokens to list.

    Returns:
        Range, median and point counts when the column has enough numbers, and
        the non-numeric tokens.
    """
    facts = []
    numbers = stats.numbers
    if len(numbers) >= privacy.MIN_VALUES:
        facts += [Fact("range", f"{privacy.format_number(min(numbers))} to "
                                f"{privacy.format_number(max(numbers))}"),
                  Fact("median", privacy.format_number(statistics.median(numbers))),
                  Fact("exact 0", str(sum(value == 0 for value in numbers))),
                  Fact("exact 100", str(sum(value == 100 for value in numbers))),
                  Fact("negative", str(sum(value < 0 for value in numbers)))]
        if stats.decimals is not None:
            facts.append(Fact("decimals", str(stats.decimals)))
    else:
        facts.append(Fact("numbers", str(len(numbers))))
    if stats.suffix:
        facts.append(Fact("suffix", repr(stats.suffix)))
    if stats.tokens:
        facts.append(Fact("tokens", counted(stats.tokens, limit)))
    return facts


def column_report(stats: ColumnStats, private: bool, limit: int, name: str | None = None,
                  levels: dict[str, str] | None = None, safe: bool = False) -> ColumnReport:
    """Report one column under the privacy rules.

    Args:
        stats: The column statistics.
        private: The table is private (rule 2).
        limit: Most categories, shapes or tokens to list.
        name: Name to print. ``None`` prints the real name.
        levels: Generic name of each category level. In a private table, a
            categorical column prints its generic level counts when given.
        safe: In a private table, the column is not a private column, so it
            prints its safe facts. Else it prints counts only.

    Returns:
        The column report.
    """
    facts = [Fact("filled", str(stats.filled)), Fact("empty", str(stats.empty)),
             Fact("unique", str(stats.unique)),
             Fact("cell types", counted(stats.cell_types, 6, quote=False))]
    report = ColumnReport(stats.name if name is None else name, stats.kind, facts)
    provider = text_provider(stats) if stats.kind == "text" else None
    if provider:
        # The name alone chooses the provider, so it prints in a private table too.
        facts.append(Fact("provider", repr(provider)))
    if private and safe:
        facts += safe_facts(stats, limit, levels)
        return report
    if private:
        if stats.kind == "numeric":
            facts.append(Fact("non-numeric", str(sum(stats.tokens.values()))))
        elif stats.kind == "categorical" and levels is not None:
            facts += safe_facts(stats, limit, levels)
        return report
    if stats.kind == "identifier":
        facts.append(identifier_shapes(stats, limit))
    elif stats.kind == "numeric":
        facts += number_facts(stats, limit)
    elif stats.kind == "boolean":
        facts += [Fact("true", str(stats.true_count)), Fact("false", str(stats.filled - stats.true_count))]
    elif stats.kind == "date":
        facts.append(Fact("format", repr(stats.date_format) if stats.date_format else "workbook date"))
        if len(stats.dates) >= privacy.MIN_VALUES:
            span = (max(stats.dates) - min(stats.dates)).total_seconds() / 86400
            facts.append(Fact("span days", privacy.format_number(span)))
    elif stats.kind == "categorical":
        shown = Counter({levels.get(value, value) if levels else value: count
                         for value, count in stats.values.items()})
        facts.append(Fact("values", counted(shown, limit)))
    elif stats.filled:
        lengths = [len(value) for value in stats.values.elements()]
        facts += [Fact("lengths", f"{min(lengths)} to {max(lengths)}"),
                  Fact("shapes", counted(Counter(privacy.shape(value) for value in stats.values.elements()),
                                         TEXT_SHAPES))]
    return report


def level_names(stats: ColumnStats) -> dict[str, str]:
    """Map each level of a categorical column to a generic name, most common first.

    A standard missing token, such as ``NA`` or ``?``, keeps its text. It holds
    no personal value, and analysis code reads it as a missing cell: the
    penguin ``Comments`` column, 84% ``NA``, had no missing cell in its
    synthetic copy when ``NA`` became ``level_1``.

    Args:
        stats: The column statistics.

    Returns:
        Each level that is not a missing level mapped to ``level_<n>``.
    """
    return privacy.generic_names([level for level, _ in stats.values.most_common()
                                  if not privacy.is_missing_level(level)], "level")


def generic_columns(columns: list[ColumnStats], generic_labels: bool) -> dict[object, str]:
    """Name the columns of a table by their kind, for ``--generic-labels``.

    Args:
        columns: The statistics of each column, in column order.
        generic_labels: Replace the column names.

    Returns:
        Each real name mapped to ``<kind>_<n>``, or no names without generic labels.
    """
    return privacy.kind_names([(stats.name, stats.kind) for stats in columns]) if generic_labels else {}


def note_columns(options: privacy.Options, file: str, table: str | None, columns: list[ColumnStats],
                 names: dict[object, str], levels: bool) -> None:
    """Add the generic column names and level names of a table to the label key.

    Args:
        options: The glimpse or profile options, which hold the key.
        file: Generic label of the file.
        table: Generic sheet name, or ``obs`` or ``var``; ``None`` for a plain table.
        columns: The statistics of each column with a generic name or generic levels.
        names: Each real column name mapped to its generic name.
        levels: The output writes generic levels for the categorical columns.
    """
    privacy.note_key(options, "column", file, table, "",
                     {stats.name: names[stats.name] for stats in columns if stats.name in names})
    if levels:
        for stats in columns:
            if stats.kind == "categorical":
                privacy.note_key(options, "level", file, table, names.get(stats.name, stats.name),
                                 level_names(stats))


#: Text that a synthetic file writes in each cell of a row above the header or
#: below the data.
NOTE_TEXT = "note"


def note_fields(row: list[object]) -> dict[str, Any]:
    """Record a row above the header or below the data, without its text.

    A title or a comment line can name a study, a person or a date, so a config
    keeps the count and the position of its cells and its comment mark only.

    Args:
        row: The row as read.

    Returns:
        The ``cells`` count, the ``start`` position when it is not 0, and the
        ``mark`` of a comment row.
    """
    filled = [position for position, cell in enumerate(row) if not is_empty(cell)]
    fields: dict[str, Any] = {"cells": len(filled)}
    if filled and filled[0]:
        fields["start"] = filled[0]
    if is_comment_row(row):
        fields["mark"] = str(row[0])[0]
    return fields


def note_cells(cells: int, start: int, mark: str | None) -> list[object]:
    """Build a row above the header or below the data of a synthetic file.

    Args:
        cells: Count of filled cells. 0 gives a blank row.
        start: Position of the first filled cell.
        mark: Comment mark at the start of the first cell, or None.

    Returns:
        Empty cells up to ``start``, then ``note`` in each filled cell.
    """
    if not cells:
        return []
    return [None] * start + [f"{mark or ''}{NOTE_TEXT}"] + [NOTE_TEXT] * (cells - 1)


def note_values(rows: list[list[object]]) -> dict[str, str]:
    """Collect the text of the rows above a header or below the data.

    Args:
        rows: The rows as read.

    Returns:
        The text of each filled cell, without its comment mark, mapped to
        ``note``, and each workbook date cell mapped to ``date``.
    """
    found: dict[str, str] = {}
    for row in rows:
        for cell in row:
            if isinstance(cell, (dt.date, dt.datetime)):
                found[cell.strftime("%Y-%m-%d")] = "date"
            elif not is_blank(cell):
                text = str(cell).lstrip("".join(COMMENT_MARKS)).strip()
                if text:
                    found[text] = "note"
    return found


def layout_facts(preamble: list[list[object]], footer: list[list[object]]) -> list[Fact]:
    """Report the rows around a table, without their text.

    Args:
        preamble: Rows above the header.
        footer: Comment rows below the data.

    Returns:
        The header row number, the comment marks and the footer row count,
        each when it applies.
    """
    facts = []
    if preamble:
        facts.append(Fact("header row", str(len(preamble) + 1)))
    marks = sorted({str(row[0])[0] for row in [*preamble, *footer] if is_comment_row(row)})
    if marks:
        facts.append(Fact("comment marks", ", ".join(repr(mark) for mark in marks)))
    if footer:
        facts.append(Fact("footer rows", str(len(footer))))
    return facts


def write_text_table(path: Path, header: list[str], rows: list[list[object]], delimiter: str,
                     preamble: list[list[object]] | None = None, footer: list[list[object]] | None = None) -> None:
    """Write a CSV or TSV file.

    Args:
        path: The output file.
        header: Column names.
        rows: Data rows of text cells.
        delimiter: The delimiter.
        preamble: Rows above the header. An empty row writes a blank line.
        footer: Rows below the data.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter=delimiter, lineterminator="\n")
        writer.writerows(preamble or [])
        writer.writerow(header)
        writer.writerows(rows)
        writer.writerows(footer or [])


def write_workbook(path: Path, sheets: list[Table]) -> None:
    """Write an Excel workbook with typed cells.

    Args:
        path: The output file.
        sheets: The tables, each with its sheet name. The rows above the header
            and the footer rows are written too.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for table in sheets:
        sheet = workbook.create_sheet(title=table.sheet)
        for row in [*table.preamble, table.header, *table.rows, *table.footer]:
            sheet.append(row)
    workbook.properties.creator = "sdg"
    workbook.save(path)
