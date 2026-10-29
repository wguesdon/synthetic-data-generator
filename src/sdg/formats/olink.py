"""Olink NPX files in the long and in the wide layout.

The long layout has one row per sample and assay. It is the CSV export of NPX
Manager and the CSV or Parquet export of Olink Explore, and it is read as CSV,
Excel or Parquet. Each column has a role, such as ``sample_id``, ``assay`` or
``value``, found from its normalized name. A column without a role, such as
``Normalization``, keeps a tabular column config and draws on its own.

The wide layout is the Excel or CSV export of NPX Manager and NPX Signature for
the qPCR platforms, as ``OlinkAnalyze`` reads it. Two head rows hold the project
name, the software version and the data type. The top rows, such as ``Panel``,
``Assay``, ``Uniprot ID`` and ``OlinkID``, describe each column. After an empty
row come the sample rows: the sample identifier, one value per assay, then the
plate, QC warning and QC deviation columns of each panel. After another empty
row, the bottom rows, such as ``LOD`` and ``Missing Data freq.``, give one cell
per assay.

Assay names, Olink IDs, UniProt IDs and panels are public, so they stay unless
``--generic-labels`` is set; sample, plate and well identifiers keep their
shape only, and the project name of a wide file becomes generic.
"""

from __future__ import annotations

import csv
import itertools
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import TYPE_CHECKING, Annotated, Any, Literal

import numpy as np
import openpyxl
import pandas as pd
from pydantic import Field, field_validator, model_validator

from sdg import privacy
from sdg.config_base import DatasetBase, Share, StrictModel
from sdg.formats import tabular
from sdg.matrix import most_common_shape, sample_names, strip_replicate
from sdg.modes.model import Feature, replicate_count, sample_count, simulate
from sdg.modes.resample import invented_identifiers, quantile_draws
from sdg.report import Fact, TableReport
from sdg.seeding import stream
from sdg.tables import (Table, analyze_column, column_report, generic_columns, is_empty, note_columns, number,
                        read_tables, read_text, sheet_records, sniff_delimiter, strip_trailing_empty, write_text_table,
                        write_workbook)

if TYPE_CHECKING:
    from sdg.generate import Context

FORMATS = ("olink",)
SUFFIXES = (".csv", ".tsv", ".txt", ".xlsx", ".parquet")

#: Role of a column, by its name in lower case without separators.
ROLES = {"sampleid": "sample_id", "index": "index", "olinkid": "olink_id", "uniprot": "uniprot",
         "assay": "assay", "panel": "panel", "plateid": "plate_id", "wellid": "well_id",
         "qcwarning": "sample_qc", "sampleqc": "sample_qc", "lod": "lod", "missingfreq": "missing_freq",
         "npx": "value", "quant": "value"}

#: Roles that a file must have to be an Olink file.
REQUIRED_ROLES = {"sample_id", "assay", "value"}

#: Value column names on the linear scale. NPX is on the log2 scale.
LINEAR_VALUES = {"quant"}

#: Well names of a 96-well plate, row by row.
WELLS = [f"{row}{column}" for row in "ABCDEFGH" for column in range(1, 13)]

#: Defaults where the privacy rules forbid a measurement.
DEFAULT_NPX = [-1.0, 0.5, 1.3, 2.0, 2.6, 3.2, 3.8, 4.5, 5.3, 6.4, 9.0]

#: Row position of the first top row of a wide file, after the two head rows.
WIDE_TOP_START = 2

#: Role of a top row of a wide file, by its normalized label.
TOP_ROLES = {"panel": "panel", "assay": "assay", "uniprotid": "uniprot", "olinkid": "olink_id", "unit": "unit"}

#: Role of a column after the assays of a wide file, by its normalized assay row text.
WIDE_COLUMN_ROLES = {"plateid": "plate_id", "qcwarning": "sample_qc", "qcdeviationfrommedian": "deviation"}

#: Role of a bottom row of a wide file, by its normalized label. Another label
#: becomes a ``number`` or a ``text`` row.
BOTTOM_ROLES = {"lod": "lod", "maxlod": "lod", "platelod": "lod", "missingdatafreq.": "missing_freq"}

#: Spread of a QC deviation column when the privacy rules forbid a measurement.
DEFAULT_DEVIATION_SD = 0.1

#: Assays per panel of the platforms whose panel size ``OlinkAnalyze`` checks
#: in a wide file, found from the panel name. Flex and Focus panels vary.
PANEL_SIZES = ((re.compile(r"target\s*96", re.IGNORECASE), "Target 96", 92),
               (re.compile(r"target\s*48", re.IGNORECASE), "Target 48", 45))

#: Normalized assay row texts of internal control columns, which ``OlinkAnalyze``
#: does not count as assays and which stay under ``--generic-labels``.
INTERNAL_CONTROLS = {"incctrl", "incctrl1", "incctrl2", "detctrl", "extctrl"}


def normalized(name: str) -> str:
    """Normalize a column name for the role lookup.

    Args:
        name: Column name.

    Returns:
        The name in lower case, without spaces, dashes or underscores.
    """
    return re.sub(r"[\s_\-]", "", name.lower())


class OlinkRoleColumn(StrictModel):
    """A column that the Olink layout fills.

    Attributes:
        name: Column header as delivered.
        kind: ``olink``.
        role: What the column holds, such as ``sample_id`` or ``value``.
    """

    name: str
    kind: Literal["olink"]
    role: Literal["sample_id", "index", "olink_id", "uniprot", "assay", "panel", "plate_id", "well_id",
                  "sample_qc", "lod", "missing_freq", "value"]


OlinkColumn = Annotated[
    tabular.IdentifierColumn | tabular.NumericColumn | tabular.CategoricalColumn | tabular.TextColumn
    | tabular.DateColumn | tabular.BooleanColumn | OlinkRoleColumn,
    Field(discriminator="kind"),
]


class OlinkAssay(StrictModel):
    """One assay of the panel.

    Attributes:
        assay: Assay name, such as ``IL6``.
        olink_id: Olink identifier, such as ``OID20563``.
        uniprot: UniProt identifier, such as ``P05231``.
        panel: Panel name.
        unit: Unit of a quantified value, which the ``Unit`` row of a wide file shows.
        lod: Limit of detection, on the value scale.
        below_lod: Share of the samples below the limit of detection.
        quantiles: Quantiles of the values across samples.
    """

    assay: str = Field(min_length=1)
    olink_id: str | None = None
    uniprot: str | None = None
    panel: str | None = None
    unit: str | None = None
    lod: float | None = None
    below_lod: Share = 0.0
    quantiles: list[float] = Field(min_length=2)

    check_quantiles = field_validator("quantiles")(tabular.check_non_decreasing)


def panel_size_error(assays: list[OlinkAssay]) -> str | None:
    """Check the assay count of each Target 96 and Target 48 panel of a wide file.

    ``OlinkAnalyze`` refuses a wide file whose Target 96 or Target 48 panels
    do not hold their full assay count, so a synthetic file must hold it too.

    Args:
        assays: The assays of the dataset.

    Returns:
        The error, or None when each panel has its count.
    """
    counts = Counter(assay.panel for assay in assays if normalized(assay.assay) not in INTERNAL_CONTROLS)
    for panel, count in counts.items():
        for pattern, platform, size in PANEL_SIZES:
            if panel and pattern.search(panel) and count != size:
                return (f"assays: the {platform} panel {panel!r} has {count} assays, but OlinkAnalyze reads "
                        f"a wide file only when each {platform} panel has its {size} assays")
    return None


class OlinkWideColumn(StrictModel):
    """One column after the assays of a wide file.

    Attributes:
        label: Text of the ``Assay`` row, such as ``Plate ID``.
        role: ``plate_id`` writes the plate of the sample, ``sample_qc`` its
            QC text, and ``deviation`` a QC deviation from the median.
        panel: Text of the ``Panel`` row.
        control: Text of the ``Uniprot ID`` row, such as ``Inc Ctrl``.
        warn_share: Share of the samples with a QC warning in a ``sample_qc``
            column. ``None`` uses ``qc_warn_share`` of the dataset.
    """

    label: str = Field(min_length=1)
    role: Literal["plate_id", "sample_qc", "deviation"]
    panel: str | None = None
    control: str | None = None
    warn_share: Share | None = None


class OlinkBottomRow(StrictModel):
    """One row under the samples of a wide file.

    Attributes:
        label: Text of the first cell, such as ``LOD``.
        role: ``lod`` writes the limit of detection of each assay,
            ``missing_freq`` its share below the limit, ``text`` one text for
            each assay, and ``number`` a value drawn from ``quantiles``.
        text: The text of a ``text`` row.
        quantiles: Quantiles of the values of a ``number`` row.
    """

    label: str = Field(min_length=1)
    role: Literal["lod", "missing_freq", "text", "number"]
    text: str = ""
    quantiles: list[float] = Field(default_factory=list)

    check_quantiles = field_validator("quantiles")(tabular.check_non_decreasing)

    @model_validator(mode="after")
    def check_number_row(self) -> OlinkBottomRow:
        """Check that a ``number`` row has quantiles.

        Returns:
            The unchanged row.

        Raises:
            ValueError: If a ``number`` row has fewer than two quantiles.
        """
        if self.role == "number" and len(self.quantiles) < 2:
            raise ValueError("quantiles: a number row needs at least two quantiles")
        return self


class OlinkWide(StrictModel):
    """The parts of a wide file around the values.

    Attributes:
        project: Text of the first cell. A profile never records the real name.
        software: Text of the second cell, such as ``NPX Signature 1.8.0``.
            ``OlinkAnalyze`` reads the version from it.
        data_type: Text of the second row, such as ``NPX Data``.
        top_rows: Labels of the rows above the samples. ``Panel``, ``Assay``,
            ``Uniprot ID``, ``OlinkID`` and ``Unit`` show a field of each assay.
        columns: The columns after the assays. ``None`` writes a ``Plate ID``
            column for each panel, then a ``QC Warning`` column for each panel.
        bottom_rows: The rows under the samples.
        deviation_sd: Spread of a QC deviation column.
    """

    project: str = "project_1"
    software: str | None = None
    data_type: str = "NPX Data"
    top_rows: list[str] = Field(default=["Panel", "Assay", "Uniprot ID", "OlinkID"], min_length=1)
    columns: list[OlinkWideColumn] | None = None
    bottom_rows: list[OlinkBottomRow] = Field(default=[OlinkBottomRow(label="LOD", role="lod"),
                                                        OlinkBottomRow(label="Missing Data freq.",
                                                                       role="missing_freq")])
    deviation_sd: float = Field(default=DEFAULT_DEVIATION_SD, ge=0.0)

    @field_validator("top_rows")
    @classmethod
    def check_top_rows(cls, labels: list[str]) -> list[str]:
        """Check that the top rows hold an ``Assay`` row, once.

        Args:
            labels: The labels of the top rows.

        Returns:
            The unchanged labels.

        Raises:
            ValueError: If no label or more than one label is ``Assay``.
        """
        if [normalized(label) for label in labels].count("assay") != 1:
            raise ValueError("a wide file needs one 'Assay' row")
        return labels


class OlinkDataset(DatasetBase):
    """An Olink file in the long or in the wide layout.

    Attributes:
        format: ``olink``.
        layout: ``long`` for one row per sample and assay, ``wide`` for the
            NPX Manager layout with one row per sample.
        delimiter: Delimiter of a text file. ``None`` uses a semicolon, as NPX Manager does.
        sheet: Sheet name of a workbook.
        columns: Columns of the long layout in delivered order: role columns
            and tabular columns. The wide layout has none.
        wide: The parts of a wide file.
        samples: Sample count in the resample mode.
        sample_pattern: Shape of a sample identifier.
        plates: Number of plates.
        plate_pattern: Shape of a plate identifier.
        qc_pass: Text of a sample that passes the QC.
        qc_warn: Text of a sample with a QC warning.
        qc_warn_share: Share of the samples with a QC warning.
        assays: The assays, in delivered order.
        scale: ``log2`` for NPX, ``linear`` for a concentration.
        decimals: Decimal places of a written value.
    """

    format: Literal["olink"]
    layout: Literal["long", "wide"] = "long"
    delimiter: str | None = Field(default=None, min_length=1, max_length=1)
    sheet: str = "Sheet1"
    columns: list[OlinkColumn] = Field(default_factory=list)
    wide: OlinkWide = Field(default_factory=OlinkWide)
    samples: int | None = Field(default=None, ge=1)
    sample_pattern: str = "SMPL-999"
    plates: int = Field(default=1, ge=1)
    plate_pattern: str = "PLT-9"
    qc_pass: str = "PASS"
    qc_warn: str = "WARN"
    qc_warn_share: Share = 0.0
    assays: list[OlinkAssay] = Field(min_length=1)
    scale: Literal["log2", "linear"] = "log2"
    decimals: int | None = Field(default=4, ge=0, le=12)

    @model_validator(mode="after")
    def check_roles(self) -> OlinkDataset:
        """Check the fields of the layout.

        A long file needs columns that hold the roles of an Olink file, once
        each. A wide file has no columns and is a text or an Excel file.

        Returns:
            The unchanged dataset.

        Raises:
            ValueError: If a field does not suit the layout, a required role
                is missing or a role repeats.
        """
        if self.layout == "wide":
            if self.columns:
                raise ValueError("columns: the wide layout has no columns; use the wide section")
            if privacy.file_suffix(Path(self.file)) == ".parquet":
                raise ValueError("file: a wide file is an .xlsx or a text file, not .parquet")
            message = panel_size_error(self.assays)
            if message:
                raise ValueError(message)
            return self
        if "wide" in self.model_fields_set:
            raise ValueError("wide: only the wide layout uses this section; set layout: wide")
        roles = [column.role for column in self.columns if column.kind == "olink"]
        missing = REQUIRED_ROLES - set(roles)
        if missing:
            raise ValueError(f"columns: an Olink file needs the roles {sorted(missing)}")
        repeated = [role for role, count in Counter(roles).items() if count > 1]
        if repeated:
            raise ValueError(f"columns: the roles {repeated} repeat")
        return self

    def feature_names(self) -> set[str]:
        """Return the assays, which effects can target.

        Returns:
            The assay names.
        """
        return {assay.assay for assay in self.assays}

    def mode_error(self, mode: str) -> str | None:
        """Require a sample count in the resample mode.

        Args:
            mode: ``resample`` or ``model``.

        Returns:
            The error, or None.
        """
        if mode == "resample" and self.samples is None:
            return "samples: the resample mode needs a sample count"
        return None


DATASET_MODELS = (OlinkDataset,)


def read_olink(path: Path) -> Table:
    """Read an Olink file of any suffix as one table.

    Args:
        path: A CSV, TSV, TXT, Excel or Parquet file.

    Returns:
        The table. Parquet cells keep their types; text cells stay text.
    """
    if privacy.file_suffix(path) == ".parquet":
        frame = pd.read_parquet(path)
        rows = [[None if (isinstance(value, float) and np.isnan(value)) else value for value in row]
                for row in frame.astype(object).itertuples(index=False, name=None)]
        return Table(None, [str(name) for name in frame.columns], rows)
    return read_tables(path)[0]


def header_of(path: Path) -> list[str]:
    """Read only the column names of a file.

    Args:
        path: The file.

    Returns:
        The column names, or an empty list for a file that does not read.
    """
    suffix = privacy.file_suffix(path)
    try:
        if suffix == ".parquet":
            import pyarrow.parquet as parquet

            return list(parquet.read_schema(path).names)
        if suffix == ".xlsx":
            workbook = openpyxl.load_workbook(path, read_only=True)
            try:
                workbook.worksheets[0].reset_dimensions()
                first = next(workbook.worksheets[0].iter_rows(max_row=1, values_only=True), ())
            finally:
                workbook.close()
            return [str(cell) for cell in first if cell is not None]
        with path.open(encoding="utf-8-sig", errors="replace") as handle:
            line = handle.readline()
        return re.split(r"[;,\t|]", line.strip())
    except (OSError, ValueError, KeyError, StopIteration):
        return []


def roles_of(header: list[str]) -> dict[str, str]:
    """Find the role of each column name.

    Args:
        header: Column names.

    Returns:
        Each role mapped to its column name. NPX wins over another value column.
    """
    found: dict[str, str] = {}
    for name in header:
        role = ROLES.get(normalized(name))
        if role and (role not in found or normalized(name) == "npx"):
            found[role] = name
    return found


def leading_cells(path: Path, count: int) -> list[str]:
    """Read the first cell of the first rows of a text or Excel file.

    Args:
        path: The file.
        count: Number of rows.

    Returns:
        The first cell of each row, or an empty list for a file that does not read.
    """
    suffix = privacy.file_suffix(path)
    try:
        if suffix == ".parquet":
            return []
        if suffix == ".xlsx":
            workbook = openpyxl.load_workbook(path, read_only=True)
            try:
                sheet = workbook.worksheets[0]
                sheet.reset_dimensions()
                rows = list(sheet.iter_rows(max_row=count, max_col=1, values_only=True))
            finally:
                workbook.close()
            return ["" if not row or row[0] is None else str(row[0]) for row in rows]
        with path.open(encoding="utf-8-sig", errors="replace") as handle:
            return [re.split(r"[;,\t|]", line.strip(), maxsplit=1)[0].strip('"')
                    for line in itertools.islice(handle, count)]
    except (OSError, ValueError, KeyError):
        return []


def is_wide(path: Path) -> bool:
    """Say whether a file is an Olink file in the wide layout.

    Args:
        path: The file.

    Returns:
        True when a ``Panel`` row and an ``Assay`` row follow the two head rows.
    """
    cells = [normalized(cell) for cell in leading_cells(path, WIDE_TOP_START + 2)]
    return cells[WIDE_TOP_START:] == ["panel", "assay"]


def sniff(path: Path) -> bool:
    """Say whether a file is an Olink file.

    Args:
        path: The file.

    Returns:
        True for a wide file, or when the header has a sample, an assay and an
        NPX or Quant column.
    """
    return is_wide(path) or REQUIRED_ROLES <= set(roles_of(header_of(path)))


@dataclass
class WideFile:
    """An Olink file in the wide layout, split into its parts.

    Attributes:
        head: The two head rows.
        top: The rows above the samples, each with its label first.
        samples: The sample rows, each with its sample identifier first.
        bottom: The rows under the samples, each with its label first.
        assay_columns: Positions of the assay columns.
        extra_columns: Position and role of each column after the assays.
        delimiter: Delimiter of a text file.
        sheet: Sheet name of a workbook.
    """

    head: list[list[object]]
    top: list[list[object]]
    samples: list[list[object]]
    bottom: list[list[object]]
    assay_columns: list[int] = field(default_factory=list)
    extra_columns: list[tuple[int, str]] = field(default_factory=list)
    delimiter: str | None = None
    sheet: str | None = None

    def top_cell(self, role: str, position: int) -> object:
        """Return the cell of a top row for one column.

        Args:
            role: Role of the top row, such as ``panel``.
            position: Column position.

        Returns:
            The cell, or None when the file has no row with that role.
        """
        for row in self.top:
            if TOP_ROLES.get(normalized(str(row[0]))) == role:
                return row[position]
        return None

    def bottom_row(self, role: str) -> list[object] | None:
        """Return the first bottom row with a role.

        Args:
            role: ``lod`` or ``missing_freq``.

        Returns:
            The row, or None.
        """
        return next((row for row in self.bottom if BOTTOM_ROLES.get(normalized(str(row[0]))) == role), None)

    def column_cells(self, position: int) -> list[object]:
        """Return the cells of one column in the sample rows.

        Args:
            position: Column position.

        Returns:
            The cells.
        """
        return [row[position] for row in self.samples]


def read_grid(path: Path) -> tuple[list[list[object]], str | None, str | None]:
    """Read every cell of a text file or of the first sheet of a workbook.

    Args:
        path: The file.

    Returns:
        The rows, all of one width with None for an empty cell, the delimiter
        of a text file and the sheet name of a workbook.
    """
    delimiter, sheet = None, None
    if privacy.file_suffix(path) == ".xlsx":
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0].title
            rows = sheet_records(workbook.worksheets[0])
        finally:
            workbook.close()
    else:
        text = read_text(path)
        delimiter = sniff_delimiter(text.split("\n", 1)[0], privacy.file_suffix(path))
        rows = [list(row) for row in csv.reader(text.splitlines(), delimiter=delimiter)]
    rows = [[None if is_empty(cell) else cell for cell in row] for row in strip_trailing_empty(rows)]
    width = max((index + 1 for row in rows for index, cell in enumerate(row) if cell is not None), default=0)
    return [(row + [None] * width)[:width] for row in rows], delimiter, sheet


def row_blocks(rows: list[list[object]]) -> list[list[list[object]]]:
    """Split rows into blocks at the empty rows.

    Args:
        rows: The rows.

    Returns:
        The runs of rows that are not empty.
    """
    blocks: list[list[list[object]]] = []
    for empty, run in itertools.groupby(rows, key=is_empty_row):
        if not empty:
            blocks.append(list(run))
    return blocks


def is_empty_row(row: list[object]) -> bool:
    """Say whether every cell of a row is empty.

    Args:
        row: The row.

    Returns:
        True for an empty row.
    """
    return all(cell is None for cell in row)


def parse_wide(path: Path) -> WideFile:
    """Read an Olink file in the wide layout.

    Args:
        path: The file.

    Returns:
        The file split into its parts, with each column classified.
    """
    rows, delimiter, sheet = read_grid(path)
    blocks = row_blocks(rows[WIDE_TOP_START:])
    wide = WideFile(rows[:WIDE_TOP_START], blocks[0] if blocks else [], blocks[1] if len(blocks) > 1 else [],
                    [row for block in blocks[2:] for row in block], delimiter=delimiter, sheet=sheet)
    width = len(rows[0]) if rows else 0
    for position in range(1, width):
        label = wide.top_cell("assay", position)
        if label is None:
            continue
        role = WIDE_COLUMN_ROLES.get(normalized(str(label)))
        if role:
            wide.extra_columns.append((position, role))
        else:
            wide.assay_columns.append(position)
    return wide


def sample_order(dataset: OlinkDataset, context: Context) -> tuple[list[str], np.ndarray]:
    """Return the sample identifiers and the values, samples by assays.

    Args:
        dataset: The dataset.
        context: The generation context, which collects the truth rows.

    Returns:
        The sample identifiers and their values.
    """
    config = context.config
    if config.mode == "model":
        features = [Feature(assay.assay, assay.quantiles, dataset.scale) for assay in dataset.assays]
        simulation = simulate(config, dataset.name, features)
        context.add_truth("effects", simulation.effects)
        samples = sample_names([(row.sample.sample_id, row.replicate) for row in simulation.rows],
                               replicate_count(config, dataset.name))
        return samples, simulation.values
    values = np.column_stack([quantile_draws(assay.quantiles,
                                             stream(config.seed, dataset.name, "assay", assay.assay).random(dataset.samples))
                              for assay in dataset.assays])
    samples = invented_identifiers(dataset.sample_pattern, dataset.samples,
                                   stream(config.seed, dataset.name, "samples"))
    return samples, values


def role_cell(role: str, sample_index: int, sample: str, assay: OlinkAssay, value: float,
              plate: str, well: str, qc: str, decimals: int | None) -> object:
    """Return the cell of a role column for one sample and assay.

    Args:
        role: The column role.
        sample_index: Position of the sample, from 0.
        sample: Sample identifier.
        assay: The assay.
        value: The value.
        plate: Plate identifier of the sample.
        well: Well of the sample.
        qc: QC text of the sample.
        decimals: Decimal places of the value.

    Returns:
        The cell.
    """
    if role == "value":
        return round(float(value), decimals) if decimals is not None else privacy.round_sig(float(value), 4)
    return {"sample_id": sample, "index": sample_index + 1, "olink_id": assay.olink_id, "uniprot": assay.uniprot,
            "assay": assay.assay, "panel": assay.panel, "plate_id": plate, "well_id": well, "sample_qc": qc,
            "lod": assay.lod, "missing_freq": assay.below_lod}[role]


def generate(dataset: OlinkDataset, context: Context) -> list[Path]:
    """Write one Olink file in the long or in the wide layout.

    Args:
        dataset: The dataset.
        context: The generation context.

    Returns:
        The file written.
    """
    seed = context.config.seed
    samples, values = sample_order(dataset, context)
    path = context.out_dir / dataset.file
    if dataset.layout == "wide":
        write_wide(dataset, wide_grid(dataset, samples, values, seed, privacy.file_suffix(path) != ".xlsx"), path)
        return [path]
    plates = invented_identifiers(dataset.plate_pattern, dataset.plates, stream(seed, dataset.name, "plates"))
    warned = stream(seed, dataset.name, "qc").random(len(samples)) < dataset.qc_warn_share
    count = len(samples) * len(dataset.assays)
    extra = {key: tabular.column_cells(column, count, stream(seed, dataset.name, key))
             for column, key in zip(dataset.columns, tabular.column_keys(dataset.columns), strict=True)
             if column.kind != "olink"}
    keys = tabular.column_keys(dataset.columns)
    rows = []
    for sample_index, sample in enumerate(samples):
        plate_index = sample_index * dataset.plates // len(samples)
        plate_start = -(-plate_index * len(samples) // dataset.plates)
        well = WELLS[(sample_index - plate_start) % len(WELLS)]
        qc = dataset.qc_warn if warned[sample_index] else dataset.qc_pass
        for assay_index, assay in enumerate(dataset.assays):
            row_index = sample_index * len(dataset.assays) + assay_index
            rows.append([role_cell(column.role, sample_index, sample, assay, values[sample_index, assay_index],
                                   plates[plate_index], well, qc, dataset.decimals) if column.kind == "olink"
                         else extra[key][row_index] for column, key in zip(dataset.columns, keys, strict=True)])
    write_olink(dataset, rows, path)
    return [path]


def write_olink(dataset: OlinkDataset, rows: list[list[object]], path: Path) -> None:
    """Write the rows in the format of the file suffix.

    Args:
        dataset: The dataset.
        rows: The rows.
        path: The output file.
    """
    header = [column.name for column in dataset.columns]
    suffix = privacy.file_suffix(path)
    if suffix == ".parquet":
        path.parent.mkdir(parents=True, exist_ok=True)
        pd.DataFrame(rows, columns=header).to_parquet(path, index=False)
    elif suffix == ".xlsx":
        write_workbook(path, [Table(dataset.sheet, header, rows)])
    else:
        delimiter = dataset.delimiter or ("\t" if suffix in {".tsv", ".txt"} else ";")
        write_text_table(path, header, [[text_cell(cell, column, dataset.decimals)
                                         for cell, column in zip(row, dataset.columns, strict=True)]
                                        for row in rows], delimiter)


def text_cell(cell: object, column: Any, decimals: int | None) -> str:
    """Write one cell of an Olink text file.

    Args:
        cell: The cell.
        column: The column config.
        decimals: Decimal places of a value.

    Returns:
        The text. A value keeps its decimal places, as ``1.2000``.
    """
    if cell is None:
        return ""
    if isinstance(cell, bool) and column.kind == "boolean":
        return column.true_text if cell else column.false_text
    if column.kind == "olink" and column.role == "value" and isinstance(cell, float) and decimals is not None:
        return f"{cell:.{decimals}f}"
    if column.kind != "olink":
        return tabular.text_cell(cell, column)
    return str(cell)


def wide_columns(dataset: OlinkDataset) -> list[OlinkWideColumn]:
    """Return the columns after the assays of a wide file.

    Args:
        dataset: The dataset.

    Returns:
        The listed columns, or by default a ``Plate ID`` column for each panel
        and then a ``QC Warning`` column for each panel.
    """
    if dataset.wide.columns is not None:
        return dataset.wide.columns
    panels = list(dict.fromkeys(assay.panel for assay in dataset.assays))
    return [OlinkWideColumn(label=label, role=role, panel=panel)
            for label, role in (("Plate ID", "plate_id"), ("QC Warning", "sample_qc")) for panel in panels]


def top_cells(label: str, assays: list[OlinkAssay], columns: list[OlinkWideColumn]) -> list[object]:
    """Return one top row of a wide file.

    Args:
        label: Label of the row, such as ``Panel``.
        assays: The assays.
        columns: The columns after the assays.

    Returns:
        The label, a cell per assay and a cell per column after the assays.
    """
    role = TOP_ROLES.get(normalized(label))
    extra = {"panel": [column.panel for column in columns], "assay": [column.label for column in columns],
             "uniprot": [column.control for column in columns]}
    return [label, *[getattr(assay, role) if role else None for assay in assays],
            *extra.get(role or "", [None] * len(columns))]


def extra_cells(dataset: OlinkDataset, columns: list[OlinkWideColumn], count: int, seed: int) -> list[list[object]]:
    """Draw the cells of the columns after the assays, one list per column.

    A plate column numbers the plates in sample order, as the long layout does,
    with plate identifiers of its own panel.

    Args:
        dataset: The dataset.
        columns: The columns after the assays.
        count: Number of samples.
        seed: The config seed.

    Returns:
        The cells of each column, one per sample.
    """
    cells = []
    for index, column in enumerate(columns):
        rng = stream(seed, dataset.name, "wide column", str(index))
        if column.role == "plate_id":
            plates = invented_identifiers(dataset.plate_pattern, dataset.plates, rng)
            cells.append([plates[sample * dataset.plates // count] for sample in range(count)])
        elif column.role == "sample_qc":
            share = dataset.qc_warn_share if column.warn_share is None else column.warn_share
            warned = rng.random(count) < share
            cells.append([dataset.qc_warn if flag else dataset.qc_pass for flag in warned])
        else:
            cells.append([round(float(value), dataset.decimals if dataset.decimals is not None else 4)
                          for value in rng.normal(0.0, dataset.wide.deviation_sd, count)])
    return cells


def bottom_cells(row: OlinkBottomRow, dataset: OlinkDataset, seed: int) -> list[object]:
    """Return one bottom row of a wide file, without the cells after the assays.

    Args:
        row: The bottom row config.
        dataset: The dataset.
        seed: The config seed.

    Returns:
        The label and a cell per assay.
    """
    assays = dataset.assays
    if row.role == "lod":
        return [row.label, *[assay.lod for assay in assays]]
    if row.role == "missing_freq":
        return [row.label, *[assay.below_lod for assay in assays]]
    if row.role == "text":
        return [row.label, *[row.text for _ in assays]]
    draws = quantile_draws(row.quantiles, stream(seed, dataset.name, "bottom row", row.label).random(len(assays)))
    return [row.label, *[privacy.round_sig(float(value), 4) for value in draws]]


def value_cell(value: float, decimals: int | None, text: bool) -> object:
    """Return one assay value of a wide file.

    Args:
        value: The value.
        decimals: Decimal places, or None for four significant digits.
        text: The file is a text file, so the value keeps its decimal places as text.

    Returns:
        The cell: None for a missing value, a float, or a text such as ``1.2000``.
    """
    if not np.isfinite(value):
        return None
    if decimals is None:
        return privacy.round_sig(float(value), 4)
    return f"{value:.{decimals}f}" if text else round(float(value), decimals)


def wide_grid(dataset: OlinkDataset, samples: list[str], values: np.ndarray, seed: int,
              text: bool) -> list[list[object]]:
    """Build every row of a wide file.

    Args:
        dataset: The dataset.
        samples: Sample identifiers.
        values: Values, samples by assays.
        seed: The config seed.
        text: The file is a text file.

    Returns:
        The rows, all of one width: the head rows, the top rows, an empty row,
        the sample rows, and an empty row before the bottom rows.
    """
    wide = dataset.wide
    columns = wide_columns(dataset)
    width = 1 + len(dataset.assays) + len(columns)
    extra = extra_cells(dataset, columns, len(samples), seed)
    rows: list[list[object]] = [[wide.project, wide.software], [wide.data_type]]
    rows += [top_cells(label, dataset.assays, columns) for label in wide.top_rows]
    rows.append([])
    for index, sample in enumerate(samples):
        rows.append([sample, *[value_cell(value, dataset.decimals, text) for value in values[index]],
                     *[cells[index] for cells in extra]])
    if wide.bottom_rows:
        rows.append([])
        rows += [bottom_cells(row, dataset, seed) for row in wide.bottom_rows]
    return [row + [None] * (width - len(row)) for row in rows]


def write_wide(dataset: OlinkDataset, rows: list[list[object]], path: Path) -> None:
    """Write the rows of a wide file in the format of the file suffix.

    Args:
        dataset: The dataset.
        rows: The rows of :func:`wide_grid`.
        path: The output file.
    """
    suffix = privacy.file_suffix(path)
    if suffix == ".xlsx":
        write_workbook(path, [Table(dataset.sheet, rows[0], rows[1:])])
        return
    delimiter = dataset.delimiter or ("\t" if suffix in {".tsv", ".txt"} else ";")
    cells = [["" if cell is None else str(cell) for cell in row] for row in rows]
    write_text_table(path, cells[0], cells[1:], delimiter)


def assay_key(row: list[object], roles: dict[str, int]) -> str:
    """Return the key that groups the rows of one assay.

    Args:
        row: A table row.
        roles: Each role mapped to its column position.

    Returns:
        The Olink identifier and the assay name.
    """
    olink_id = row[roles["olink_id"]] if "olink_id" in roles else ""
    return f"{olink_id}|{row[roles['assay']]}"


def assay_configs(table: Table, roles: dict[str, int], private: bool,
                  options: privacy.Options) -> tuple[list[dict[str, Any]], list[str]]:
    """Profile each assay: its public names, its limit and its value quantiles.

    Args:
        table: The table.
        roles: Each role mapped to its column position.
        private: The table records counts only.
        options: The profile options.

    Returns:
        The assay configs and the paths of their default fields.
    """
    groups: dict[str, list[list[object]]] = {}
    for row in table.rows:
        groups.setdefault(assay_key(row, roles), []).append(row)
    assays, defaults = [], []
    for index, rows in enumerate(groups.values()):
        first = rows[0]
        values = [value for row in rows if (value := number(row[roles["value"]])) is not None]
        lod = number(first[roles["lod"]]) if "lod" in roles else None
        name = f"assay_{index + 1}" if options.generic_labels else str(first[roles["assay"]])
        assay: dict[str, Any] = {"assay": name}
        for role in ("olink_id", "uniprot", "panel"):
            if role in roles and not options.generic_labels:
                assay[role] = str(first[roles[role]])
        if lod is not None:
            assay["lod"] = privacy.round_sig(lod)
            assay["below_lod"] = tabular.share(sum(value < lod for value in values), len(values))
        if private or len(values) < privacy.MIN_VALUES:
            assay["quantiles"] = list(DEFAULT_NPX)
            defaults.append(f"assays[{index}].quantiles")
        else:
            assay["quantiles"] = tabular.rounded_quantiles(values)
        assays.append(assay)
    return assays, defaults


def column_configs(table: Table, roles: dict[str, int], private: bool, options: privacy.Options,
                   label: str) -> tuple[list[dict[str, Any]], list[str]]:
    """Profile the columns: a role column keeps its role, another its tabular config.

    A role column keeps its standard name, such as ``SampleID`` or ``NPX``,
    because the Olink reader and the analysis code find the roles by name. With
    generic labels, the other columns get the generic name of their kind.

    Args:
        table: The table.
        roles: Each role mapped to its column position.
        private: The table records counts only.
        options: The profile options.
        label: Generic label of the file, for the label key.

    Returns:
        The column configs and the paths of their default fields.
    """
    by_position = {position: role for role, position in roles.items()}
    # With generic labels, a role column counts in the numbering of its kind,
    # so a generic name is the same in the profile and in the glimpse.
    analyzed = {index: analyze_column(name, table.column(index), options.max_categories)
                for index, name in enumerate(table.header) if options.generic_labels or index not in by_position}
    names = generic_columns(list(analyzed.values()), options.generic_labels)
    others = [stats for index, stats in analyzed.items() if index not in by_position]
    sheet = "sheet_1" if options.generic_labels and table.sheet else table.sheet
    note_columns(options, label, sheet, others, names, private or options.generic_labels)
    columns, defaults = [], []
    for index, name in enumerate(table.header):
        if index in by_position:
            columns.append({"name": name, "kind": "olink", "role": by_position[index]})
            continue
        stats = analyzed[index]
        column, defaulted = tabular.column_fields(stats, names.get(name, name), private, options.generic_labels,
                                                  False)
        columns.append(column)
        defaults += [f"columns[{index}].{item}" for item in defaulted]
    return columns, defaults


def note_sheet(options: privacy.Options, label: str, sheet: str | None) -> None:
    """Add the generic name of the one sheet of a workbook to the label key.

    Args:
        options: The glimpse or profile options.
        label: Generic label of the file.
        sheet: The real sheet name, or None for a text file.
    """
    if options.generic_labels and sheet:
        privacy.note_key(options, "sheet", label, None, "", {sheet: "sheet_1"})


def profile_file(path: Path, label: str, options: privacy.Options) -> tuple[dict[str, Any], list[str]]:
    """Profile one Olink file.

    Args:
        path: The file.
        label: Generic file label.
        options: The profile options.

    Returns:
        The dataset config and its default fields.
    """
    table = read_olink(path)
    names = roles_of(table.header)
    roles = {role: table.header.index(name) for role, name in names.items()}
    private = tabular.is_private_table(path, table, options)
    samples = list(dict.fromkeys(str(row[roles["sample_id"]]) for row in table.rows))
    dataset: dict[str, Any] = {"name": label.split(".", 1)[0], "format": "olink", "file": label}
    if table.delimiter and table.delimiter != ";":
        dataset["delimiter"] = table.delimiter
    if table.sheet:
        dataset["sheet"] = "sheet_1" if options.generic_labels else table.sheet
        note_sheet(options, label, table.sheet)
    columns, defaults = column_configs(table, roles, private, options, label)
    dataset.update({"columns": columns, "samples": len(samples), "sample_pattern": most_common_shape(samples)})
    if "plate_id" in roles:
        plates = list(dict.fromkeys(str(row[roles["plate_id"]]) for row in table.rows))
        dataset.update({"plates": len(plates), "plate_pattern": most_common_shape(plates)})
    if "sample_qc" in roles:
        by_sample = {str(row[roles["sample_id"]]): str(row[roles["sample_qc"]]) for row in table.rows}
        qc = Counter(by_sample.values()).most_common()
        dataset["qc_pass"] = qc[0][0]
        if len(qc) > 1:
            dataset.update({"qc_warn": qc[1][0],
                            "qc_warn_share": tabular.share(len(by_sample) - qc[0][1], len(by_sample))})
    assays, assay_defaults = assay_configs(table, roles, private, options)
    value_name = names["value"]
    places = [place for cell in table.column(roles["value"])[:500]
              if isinstance(cell, str) and (place := decimals_of_text(cell)) is not None]
    dataset.update({"assays": assays, "scale": "linear" if normalized(value_name) in LINEAR_VALUES else "log2",
                    "decimals": max(places) if places else 4})
    return dataset, defaults + assay_defaults


def decimals_of_text(cell: str) -> int | None:
    """Count the decimal places of a number written as text.

    Args:
        cell: The text.

    Returns:
        The decimal places, or None for a text that is not a plain number.
    """
    if number(cell) is None or "e" in cell.lower():
        return None
    return len(cell.strip().split(".", 1)[1]) if "." in cell else 0


def profile(paths: list[Path], labels: list[str], options: privacy.Options) -> list[tuple[dict[str, Any], list[str]]]:
    """Profile Olink files.

    Args:
        paths: The files.
        labels: Generic label of each file.
        options: The profile options.

    Returns:
        One dataset config per file, with the paths of its default fields.
    """
    return [profile_wide(path, label, options) if is_wide(path) else profile_file(path, label, options)
            for path, label in zip(paths, labels, strict=True)]


def panel_names(wide: WideFile, options: privacy.Options) -> dict[str, str]:
    """Return the name to record for each panel of a wide file.

    Args:
        wide: The wide file.
        options: The profile options.

    Returns:
        Each real panel name mapped to itself, or to ``panel_<n>`` with
        ``--generic-labels``. A wide file needs the panels for its layout.
    """
    panels = [str(cell) for position in wide.assay_columns + [item[0] for item in wide.extra_columns]
              if (cell := wide.top_cell("panel", position)) is not None]
    unique = list(dict.fromkeys(panels))
    if options.generic_labels:
        return {panel: f"panel_{index}" for index, panel in enumerate(unique, start=1)}
    return {panel: panel for panel in unique}


def below_lod_share(values: list[float], lod: float | None, frequency: object) -> float | None:
    """Return the share of values below the limit of detection.

    Args:
        values: The values of one assay.
        lod: The limit of detection, or None.
        frequency: The cell of the ``Missing Data freq.`` row, or None.

    Returns:
        The share from the frequency cell, or else from the values and the
        limit, or None when neither is known.
    """
    stated = number(frequency)
    if stated is not None and 0.0 <= stated <= 1.0:
        return privacy.round_sig(stated)
    if lod is not None and values:
        return tabular.share(sum(value < lod for value in values), len(values))
    return None


def wide_assay_configs(wide: WideFile, private: bool, panels: dict[str, str],
                       options: privacy.Options) -> tuple[list[dict[str, Any]], list[str]]:
    """Profile each assay column of a wide file.

    Args:
        wide: The wide file.
        private: The file records counts only.
        panels: The name to record for each panel.
        options: The profile options.

    Returns:
        The assay configs and the paths of their default fields.
    """
    lod_row, frequency_row = wide.bottom_row("lod"), wide.bottom_row("missing_freq")
    assays, defaults = [], []
    for index, position in enumerate(wide.assay_columns):
        values = [value for cell in wide.column_cells(position) if (value := number(cell)) is not None]
        real = str(wide.top_cell("assay", position))
        generic = options.generic_labels and normalized(real) not in INTERNAL_CONTROLS
        assay: dict[str, Any] = {"assay": f"assay_{index + 1}" if generic else real}
        for role in ("olink_id", "uniprot", "unit"):
            cell = wide.top_cell(role, position)
            if cell is not None and not options.generic_labels:
                assay[role] = str(cell)
        panel = wide.top_cell("panel", position)
        if panel is not None:
            assay["panel"] = panels[str(panel)]
        lod = number(lod_row[position]) if lod_row else None
        if lod is not None:
            assay["lod"] = privacy.round_sig(lod)
        below = below_lod_share(values, lod, frequency_row[position] if frequency_row else None)
        if below is not None:
            assay["below_lod"] = below
        if private or len(values) < privacy.MIN_VALUES:
            assay["quantiles"] = list(DEFAULT_NPX)
            defaults.append(f"assays[{index}].quantiles")
        else:
            assay["quantiles"] = tabular.rounded_quantiles(values)
        assays.append(assay)
    return assays, defaults


def wide_bottom_configs(wide: WideFile, private: bool) -> tuple[list[dict[str, Any]], list[str]]:
    """Profile the bottom rows of a wide file.

    A row with a known label keeps its role. Another row is a ``number`` row
    when each of its assay cells is a number, or else a ``text`` row with its
    most common text when that text repeats.

    Args:
        wide: The wide file.
        private: The file records counts only.

    Returns:
        The bottom row configs and the paths of their default fields.
    """
    rows, defaults = [], []
    for index, row in enumerate(wide.bottom):
        label = str(row[0])
        role = BOTTOM_ROLES.get(normalized(label))
        if role:
            rows.append({"label": label, "role": role})
            continue
        cells = [row[position] for position in wide.assay_columns if row[position] is not None]
        numbers = [value for cell in cells if (value := number(cell)) is not None]
        if cells and len(numbers) == len(cells):
            if private or len(numbers) < privacy.MIN_VALUES:
                rows.append({"label": label, "role": "number", "quantiles": [0.0, 1.0]})
                defaults.append(f"wide.bottom_rows[{index}].quantiles")
            else:
                rows.append({"label": label, "role": "number", "quantiles": tabular.rounded_quantiles(numbers)})
            continue
        common = Counter(str(cell) for cell in cells).most_common(1)
        rows.append({"label": label, "role": "text",
                     "text": common[0][0] if common and common[0][1] > 1 else "value"})
    return rows, defaults


def wide_extra_configs(wide: WideFile, panels: dict[str, str], qc_pass: str | None) -> list[dict[str, Any]]:
    """Profile the columns after the assays of a wide file.

    A QC column records its own warning share, because a rare warning can be
    absent from one panel; a share pooled over the panels would then put a
    warning into that panel.

    Args:
        wide: The wide file.
        panels: The name to record for each panel.
        qc_pass: The QC text of a sample that passes, or None.

    Returns:
        The column configs.
    """
    columns = []
    for position, role in wide.extra_columns:
        column: dict[str, Any] = {"label": str(wide.top_cell("assay", position)), "role": role}
        panel = wide.top_cell("panel", position)
        if panel is not None:
            column["panel"] = panels[str(panel)]
        control = wide.top_cell("uniprot", position)
        if role == "deviation" and control is not None:
            column["control"] = str(control)
        cells = [str(cell) for cell in wide.column_cells(position) if cell is not None]
        if role == "sample_qc" and qc_pass is not None and cells:
            column["warn_share"] = tabular.share(sum(cell != qc_pass for cell in cells), len(cells))
        columns.append(column)
    return columns


def wide_cells(wide: WideFile, role: str) -> list[object]:
    """Return the filled sample cells of every column after the assays with one role.

    Args:
        wide: The wide file.
        role: ``plate_id``, ``sample_qc`` or ``deviation``.

    Returns:
        The cells, column by column.
    """
    return [cell for position, found in wide.extra_columns if found == role
            for cell in wide.column_cells(position) if cell is not None]


def profile_wide(path: Path, label: str, options: privacy.Options) -> tuple[dict[str, Any], list[str]]:
    """Profile one Olink file in the wide layout.

    Args:
        path: The file.
        label: Generic file label.
        options: The profile options.

    Returns:
        The dataset config and its default fields.
    """
    wide = parse_wide(path)
    private = bool(options.private_files.search(path.name))
    panels = panel_names(wide, options)
    samples = [str(row[0]) for row in wide.samples]
    dataset: dict[str, Any] = {"name": label.split(".", 1)[0], "format": "olink", "file": label, "layout": "wide"}
    if wide.delimiter and wide.delimiter != ";":
        dataset["delimiter"] = wide.delimiter
    if wide.sheet:
        dataset["sheet"] = "sheet_1" if options.generic_labels else wide.sheet
        note_sheet(options, label, wide.sheet)
    dataset.update({"samples": len(samples), "sample_pattern": most_common_shape(samples)})
    plate_columns = [position for position, role in wide.extra_columns if role == "plate_id"]
    if plate_columns:
        plates = [str(cell) for cell in wide_cells(wide, "plate_id")]
        dataset.update({"plates": max(len({cell for cell in wide.column_cells(position) if cell is not None})
                                      for position in plate_columns) or 1,
                        "plate_pattern": most_common_shape(plates)})
    qc = Counter(str(cell) for cell in wide_cells(wide, "sample_qc")).most_common()
    if qc:
        dataset["qc_pass"] = qc[0][0]
        if len(qc) > 1:
            total = sum(count for _, count in qc)
            dataset.update({"qc_warn": qc[1][0], "qc_warn_share": tabular.share(total - qc[0][1], total)})
    assays, defaults = wide_assay_configs(wide, private, panels, options)
    bottom, bottom_defaults = wide_bottom_configs(wide, private)
    data_type = str(wide.head[1][0]) if len(wide.head) > 1 and wide.head[1][0] is not None else "NPX Data"
    places = [place for position in wide.assay_columns[:50] for cell in wide.column_cells(position)
              if isinstance(cell, str) and (place := decimals_of_text(cell)) is not None]
    section: dict[str, Any] = {"data_type": data_type, "top_rows": [str(row[0]) for row in wide.top],
                               "columns": wide_extra_configs(wide, panels, qc[0][0] if qc else None),
                               "bottom_rows": bottom}
    if wide.head and len(wide.head[0]) > 1 and wide.head[0][1] is not None:
        section["software"] = str(wide.head[0][1])
    deviations = [value for cell in wide_cells(wide, "deviation") if (value := number(cell)) is not None]
    if len(deviations) >= privacy.MIN_VALUES and not private:
        section["deviation_sd"] = privacy.round_sig(float(np.std(deviations)))
    dataset.update({"assays": assays, "scale": "linear" if "quant" in data_type.lower() else "log2",
                    "decimals": max(places) if places else 4, "wide": section})
    return dataset, defaults + bottom_defaults


def glimpse(paths: list[Path], labels: list[str], options: privacy.Options) -> list[TableReport]:
    """Report Olink files: the tabular report with the Olink facts first.

    Args:
        paths: The files.
        labels: Generic label of each file.
        options: The glimpse options.

    Returns:
        One table report per file.
    """
    reports = []
    for path, label in zip(paths, labels, strict=True):
        if is_wide(path):
            reports.append(glimpse_wide(path, label, options))
            continue
        table = read_olink(path)
        names = roles_of(table.header)
        roles = {role: table.header.index(name) for role, name in names.items()}
        sheet = "sheet_1" if options.generic_labels and table.sheet else table.sheet
        note_sheet(options, label, table.sheet)
        # A role column keeps its standard name, as in the profile, because the
        # Olink reader and the analysis code find the roles by name.
        report = tabular.glimpse_table(path, label, table, sheet, options, frozenset(roles.values()))
        report.format = "olink"
        samples = {str(row[roles["sample_id"]]) for row in table.rows}
        assays = {assay_key(row, roles) for row in table.rows}
        facts = [Fact("samples", str(len(samples))), Fact("assays", str(len(assays))),
                 Fact("value column", repr(names["value"]))]
        if "panel" in roles and not options.generic_labels:
            panels = sorted({str(row[roles["panel"]]) for row in table.rows})
            facts.append(Fact("panels", ", ".join(repr(panel) for panel in panels)))
        report.facts = facts + report.facts
        reports.append(report)
    return reports


def glimpse_wide(path: Path, label: str, options: privacy.Options) -> TableReport:
    """Report one Olink file in the wide layout.

    The report lists the labels of the top and bottom rows, then one column
    for the sample identifiers, one for all assay values, and one for each
    column after the assays.

    Args:
        path: The file.
        label: Generic label of the file.
        options: The glimpse options.

    Returns:
        The table report.
    """
    wide = parse_wide(path)
    private = bool(options.private_files.search(path.name))
    width = len(wide.top[0]) if wide.top else 0
    data_type = wide.head[1][0] if len(wide.head) > 1 else None
    software = wide.head[0][1] if wide.head and len(wide.head[0]) > 1 else None
    facts = [Fact("layout", "wide"), Fact("samples", str(len(wide.samples))),
             Fact("assays", str(len(wide.assay_columns))), Fact("columns", str(width)),
             Fact("data type", repr("" if data_type is None else str(data_type))),
             Fact("software", repr("" if software is None else str(software))),
             Fact("top rows", ", ".join(repr(str(row[0])) for row in wide.top)),
             Fact("bottom rows", ", ".join(repr(str(row[0])) for row in wide.bottom) or "none")]
    if not options.generic_labels:
        panels = list(dict.fromkeys(str(cell) for position in wide.assay_columns
                                    if (cell := wide.top_cell("panel", position)) is not None))
        facts.append(Fact("panels", ", ".join(repr(panel) for panel in sorted(panels))))
    if wide.delimiter:
        facts.append(Fact("delimiter", repr(wide.delimiter)))
    if private:
        facts.append(Fact("private", "counts only"))
    sheet = "sheet_1" if options.generic_labels and wide.sheet else wide.sheet
    note_sheet(options, label, wide.sheet)
    report = TableReport(tabular.table_label(label, sheet), "olink", facts)
    limit = options.max_categories
    stats = analyze_column("sample", [row[0] for row in wide.samples], limit)
    report.columns.append(column_report(stats, private, limit))
    values = [cell for position in wide.assay_columns for cell in wide.column_cells(position)]
    report.columns.append(column_report(analyze_column("assay values", values, limit), private, limit))
    for position, _ in wide.extra_columns:
        stats = analyze_column(str(wide.top_cell("assay", position)), wide.column_cells(position), limit)
        report.columns.append(column_report(stats, private, limit))
    return report


def audit_values(paths: list[Path], options: privacy.Options) -> dict[str, str]:
    """Collect the sample, plate and well identifiers and the other sensitive cells.

    Args:
        paths: The real files.
        options: The options that made the checked text.

    Returns:
        Each sensitive value mapped to its source kind.
    """
    found: dict[str, str] = {}
    for path in paths:
        if is_wide(path):
            found.update(wide_audit_values(path, options))
            continue
        table = read_olink(path)
        roles = {role: table.header.index(name) for role, name in roles_of(table.header).items()}
        for role in ("sample_id", "plate_id", "well_id"):
            if role in roles:
                found.update({str(cell): role.replace("_", " ") for cell in table.column(roles[role])
                              if cell is not None and len(str(cell)) > 3})
        if options.generic_labels and table.sheet:
            found[table.sheet] = "sheet name"
        for index, name in enumerate(table.header):
            if index in roles.values():
                continue
            if options.generic_labels:
                found[str(name)] = "column name"
            stats = analyze_column(name, table.column(index), options.max_categories)
            if stats.kind in tabular.AUDITED_KINDS:
                found.update({value: tabular.AUDITED_KINDS[stats.kind] for value in stats.values})
    return found


def wide_audit_values(path: Path, options: privacy.Options) -> dict[str, str]:
    """Collect the sensitive cells of a wide file.

    Args:
        path: The real file.
        options: The options that made the checked text.

    Returns:
        The sample and plate identifiers and the project name, and with
        ``--generic-labels`` the assay, panel and sheet names too.
    """
    wide = parse_wide(path)
    found = {str(row[0]): "sample id" for row in wide.samples if row[0] is not None and len(str(row[0])) > 3}
    found.update({str(cell): "plate id" for cell in wide_cells(wide, "plate_id") if len(str(cell)) > 3})
    if wide.head and wide.head[0][0] is not None:
        found[str(wide.head[0][0])] = "project name"
    if options.generic_labels:
        for role in ("assay", "panel"):
            found.update({str(cell): f"{role} name" for position in wide.assay_columns
                          if (cell := wide.top_cell(role, position)) is not None})
        if wide.sheet:
            found[wide.sheet] = "sheet name"
    return found


TEMPLATE = '''\
  # An Olink file in the long layout: one row per sample and assay. The model
  # mode takes the samples of the design; the resample mode needs samples.
  # For the NPX Manager wide file, set layout: wide, remove columns and add a
  # wide section (see docs/how_it_works.md).
  - name: {name}
    format: olink
    layout: long                      # long, or wide (.xlsx or .csv only)
    file: {name}.csv                  # .csv (semicolon), .xlsx or .parquet
    columns:                          # delivered order; kind olink fills a role
      - {{name: SampleID, kind: olink, role: sample_id}}
      - {{name: Index, kind: olink, role: index}}
      - {{name: OlinkID, kind: olink, role: olink_id}}
      - {{name: UniProt, kind: olink, role: uniprot}}
      - {{name: Assay, kind: olink, role: assay}}
      - {{name: MissingFreq, kind: olink, role: missing_freq}}
      - {{name: Panel, kind: olink, role: panel}}
      - {{name: PlateID, kind: olink, role: plate_id}}
      - {{name: QC_Warning, kind: olink, role: sample_qc}}
      - {{name: LOD, kind: olink, role: lod}}
      - {{name: NPX, kind: olink, role: value}}
      - {{name: Normalization, kind: categorical, levels: {{Intensity: 1.0}}}}
{samples}    plates: 1
    qc_warn_share: 0.02
    scale: log2                       # NPX is log2; a Quant column is linear
    decimals: 4
    assays:                           # public names; quantiles of the values across samples
      - {{assay: IL6, olink_id: OID20563, uniprot: P05231, panel: Olink Target 96 Inflammation,
          lod: 1.2, below_lod: 0.1, quantiles: [0.5, 1.5, 2.2, 2.8, 3.3, 3.8, 4.3, 4.9, 5.6, 6.6, 9.5]}}
      - {{assay: TNF, olink_id: OID20581, uniprot: P01375, panel: Olink Target 96 Inflammation,
          lod: 0.8, below_lod: 0.05, quantiles: [0.6, 1.4, 1.8, 2.1, 2.4, 2.7, 3.0, 3.3, 3.7, 4.3, 6.0]}}
'''


def template(fmt: str, mode: str, name: str) -> tuple[str, str | None]:
    """Return the commented template of an Olink dataset.

    Args:
        fmt: ``olink``.
        mode: ``resample`` or ``model``.
        name: Dataset name.

    Returns:
        The YAML list item, and the feature that a template effect can target.
    """
    samples = "" if mode == "model" else "    samples: 40\n    sample_pattern: SMPL-999  # A is a letter, 9 a digit; other characters stay\n"
    return TEMPLATE.format(name=name, samples=samples), "IL6" if mode == "model" else None


def check_output(dataset: OlinkDataset, config: Any, out_dir: Path) -> list[str]:
    """Compare the output file with the config: header and rows, or the wide parts.

    Args:
        dataset: The dataset config.
        config: The whole config.
        out_dir: The output folder.

    Returns:
        The problems found.
    """
    path = out_dir / dataset.file
    if not path.is_file():
        return [f"{dataset.file}: missing"]
    if dataset.layout == "wide":
        return check_wide(dataset, parse_wide(path), sample_count(dataset, config))
    table = read_olink(path)
    header = [column.name for column in dataset.columns]
    expected = sample_count(dataset, config) * len(dataset.assays)
    issues = []
    if table.header != header:
        issues.append(f"{dataset.file}: header {table.header} differs from the config {header}")
    if len(table.rows) != expected:
        issues.append(f"{dataset.file}: {len(table.rows)} rows, the config gives {expected}")
    return issues


def observed(dataset: OlinkDataset, out_dir: Path) -> dict[tuple[str, str], list[float]]:
    """Read the values back, by sample and assay.

    Args:
        dataset: The dataset config.
        out_dir: The output folder.

    Returns:
        The values of each sample and assay.
    """
    if dataset.layout == "wide":
        return observed_wide(dataset, parse_wide(out_dir / dataset.file))
    table = read_olink(out_dir / dataset.file)
    roles = {column.role: index for index, column in enumerate(dataset.columns) if column.kind == "olink"}
    found: dict[tuple[str, str], list[float]] = {}
    for row in table.rows:
        value = number(row[roles["value"]])
        if value is not None:
            key = (strip_replicate(str(row[roles["sample_id"]])), str(row[roles["assay"]]))
            found.setdefault(key, []).append(value)
    return found


def check_wide(dataset: OlinkDataset, wide: WideFile, expected: int) -> list[str]:
    """Compare a wide output file with the config.

    Args:
        dataset: The dataset config.
        wide: The output file as read.
        expected: Number of samples that the config gives.

    Returns:
        The problems found.
    """
    issues = []
    checks = [("samples", len(wide.samples), expected),
              ("assay columns", len(wide.assay_columns), len(dataset.assays)),
              ("columns after the assays", len(wide.extra_columns), len(wide_columns(dataset))),
              ("top rows", [str(row[0]) for row in wide.top], dataset.wide.top_rows),
              ("bottom rows", [str(row[0]) for row in wide.bottom], [row.label for row in dataset.wide.bottom_rows])]
    for name, found, wanted in checks:
        if found != wanted:
            issues.append(f"{dataset.file}: {name} {found}, the config gives {wanted}")
    return issues


def observed_wide(dataset: OlinkDataset, wide: WideFile) -> dict[tuple[str, str], list[float]]:
    """Read the values of a wide file back, by sample and assay.

    Args:
        dataset: The dataset config.
        wide: The output file as read.

    Returns:
        The values of each sample and assay.
    """
    found: dict[tuple[str, str], list[float]] = {}
    for row in wide.samples:
        for assay, position in zip(dataset.assays, wide.assay_columns, strict=True):
            value = number(row[position])
            if value is not None:
                found.setdefault((strip_replicate(str(row[0])), assay.assay), []).append(value)
    return found
